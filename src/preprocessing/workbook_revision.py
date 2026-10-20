"""Workbook de revisión manual: estructura común a todas las fuentes (MET, CMA y AIC).

Cada fuente tiene su generador (src/preprocessing/create_<fuente>_review_workbook_v2.py),
que solo traduce sus registros a las columnas de COLUMNAS; este módulo arma el workbook:

- hoja «resumen»: registros por sugerencia automática (corpus_actual) e instrucciones;
- hoja de revisión (nombre en src/utils/fuentes.py): una fila por registro, con listas
  desplegables en corpus_final y decision_revision e hipervínculos en las URL;
- hoja «listas»: los valores admitidos.

El workbook es el registro de las decisiones humanas: si ya existe, no se sobrescribe
salvo con --forzar.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.utils.fuentes import DECISIONES, FUENTES, SUBCONJUNTOS

# Columna y ancho en Excel. Las ocho primeras registran la decision; el resto describe el registro.
COLUMNAS_ANCHOS = [
    ("id_fuente", 14), ("corpus_actual", 16), ("corpus_final", 16), ("decision_revision", 18),
    ("motivo_revision", 42), ("observaciones_revision", 46), ("revisor", 18), ("fecha_revision", 18),
    ("titulo_original", 34), ("titulo_es_sugerido", 34), ("cultura", 20), ("fecha_objeto", 20),
    ("procedencia", 24), ("tipo_objeto", 20), ("tipo_superficie", 24), ("material", 34),
    ("material_normalizado", 28), ("tecnica", 18), ("estado_curacion", 18), ("destino_curacion", 20),
    ("motivo_curacion", 58), ("url_objeto", 42), ("url_imagen", 42), ("ruta_imagen_local", 36),
    ("licencia", 34), ("origen_curacion", 24),
]
COLUMNAS = [columna for columna, _ in COLUMNAS_ANCHOS]
OPCIONES_DECISION = ["mantener", "mover", "descartar", "recuperar", "revisar"]
INSTRUCCIONES = [
    "corpus_final: principal, secundario, descartados o revisar. Mientras diga revisar, el registro queda pendiente.",
    "corpus_actual es la sugerencia automatica o la clasificacion de partida; no se modifica.",
    "decision_revision, motivo_revision, revisor y fecha_revision documentan la decision humana.",
    "Las filas con revisor = auto son decisiones automaticas documentadas; el revisor puede corregirlas.",
]


def crear_hoja_resumen(libro: Workbook, titulo: str, filas: list[dict[str, str]], notas: list[str]) -> None:
    hoja = libro.create_sheet("resumen", 0)
    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A3"], hoja["B3"] = "Sugerencia automatica (corpus_actual)", "Registros"
    hoja["A3"].font = hoja["B3"].font = Font(bold=True)
    conteo = Counter(f["corpus_actual"] for f in filas)
    for numero, corpus in enumerate(SUBCONJUNTOS, start=4):
        hoja[f"A{numero}"], hoja[f"B{numero}"] = corpus, conteo.get(corpus, 0)
    hoja["A9"] = "Uso"
    hoja["A9"].font = Font(bold=True)
    for numero, texto in enumerate(INSTRUCCIONES + notas, start=10):
        hoja[f"A{numero}"] = texto
    hoja.column_dimensions["A"].width = 110
    hoja.column_dimensions["B"].width = 14


def crear_hoja_revision(libro: Workbook, nombre: str, filas: list[dict[str, str]]) -> None:
    hoja = libro.create_sheet(nombre)
    hoja.append(COLUMNAS)
    for fila in filas:
        hoja.append([fila.get(columna, "") for columna in COLUMNAS])

    for celda in hoja[1]:
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if COLUMNAS[celda.column - 1] in ("url_objeto", "url_imagen") and celda.value:
                celda.hyperlink, celda.style = celda.value, "Hyperlink"
    for numero, (_, ancho) in enumerate(COLUMNAS_ANCHOS, start=1):
        hoja.column_dimensions[get_column_letter(numero)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    ultima = max(hoja.max_row, 2)
    for columna, opciones in (("corpus_final", DECISIONES), ("decision_revision", OPCIONES_DECISION)):
        validacion = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=False)
        hoja.add_data_validation(validacion)
        letra = get_column_letter(COLUMNAS.index(columna) + 1)
        validacion.add(f"{letra}2:{letra}{ultima}")


def crear_hoja_listas(libro: Workbook) -> None:
    hoja = libro.create_sheet("listas")
    for numero, (nombre, opciones) in enumerate((("corpus_final", DECISIONES), ("decision_revision", OPCIONES_DECISION)), start=1):
        letra = get_column_letter(numero)
        hoja[f"{letra}1"] = nombre
        hoja[f"{letra}1"].font = Font(bold=True)
        hoja[f"{letra}1"].fill = PatternFill("solid", fgColor="D9EAF7")
        for fila, opcion in enumerate(opciones, start=2):
            hoja[f"{letra}{fila}"] = opcion
        hoja.column_dimensions[letra].width = 24


def crear_workbook(raiz: Path, fuente: str, construir_filas, notas: list[str], forzar: bool) -> int:
    """Escribe el workbook de la fuente; no sobrescribe uno existente sin forzar."""
    registro = FUENTES[fuente]
    salida = raiz / registro.workbook
    if salida.exists() and not forzar:
        print(f"El workbook ya existe y contiene decisiones manuales: {salida}")
        print("No se sobrescribio. Use --forzar solo si quiere regenerarlo desde cero.")
        return 1
    filas = construir_filas(raiz)
    libro = Workbook()
    libro.remove(libro.active)
    crear_hoja_resumen(libro, f"Revision manual {fuente} v2", filas, notas)
    crear_hoja_revision(libro, registro.hoja, filas)
    crear_hoja_listas(libro)
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    print(f"Workbook creado: {salida}")
    print(f"Registros para revision: {len(filas)}")
    return 0


def main_fuente(argv: list[str] | None, fuente: str, construir_filas, notas: list[str]) -> int:
    """Punto de entrada común de los generadores por fuente."""
    parser = argparse.ArgumentParser(description=f"Crea el workbook maestro de revision manual {fuente} v2.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    parser.add_argument("--forzar", action="store_true",
                        help="Sobrescribe el workbook existente y pierde las decisiones manuales registradas.")
    args = parser.parse_args(argv)
    return crear_workbook(Path(args.root).resolve(), fuente, construir_filas, notas, args.forzar)
