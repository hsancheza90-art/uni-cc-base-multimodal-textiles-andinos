"""Crea el workbook maestro de revision manual AIC v2.

Uso: python -m src.preprocessing.create_aic_review_workbook_v2 [--forzar]

Mismas columnas, listas y validaciones que el workbook CMA v2. Cada candidato
llega con la sugerencia automatica en corpus_actual (principal, secundario o
descartados) y corpus_final = revisar, salvo las herramientas textiles, que se
descartan automaticamente como en CMA (revisor "auto"). La decision humana se
registra en corpus_final, decision_revision, motivo_revision, revisor y fecha.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from src.preprocessing.create_cma_review_workbook_v2 import (
    COLUMNAS,
    agregar_hipervinculos,
    agregar_validaciones,
    aplicar_estilo_revision,
    crear_hoja_listas,
)
from src.utils.archivos import leer_csv

CANDIDATOS = "data/metadata/aic_andes_textiles_candidates.csv"
WORKBOOK = "data/metadata/aic_revision_manual_v2.xlsx"
HOJA_REVISION = "revision_aic_v2"


def es_herramienta(candidato: dict[str, str]) -> bool:
    return candidato["estado_filtro"] == "excluido_no_textil" and "herramienta textil" in candidato["motivo_filtro"]


def construir_filas(raiz: Path) -> list[dict[str, str]]:
    filas = []
    for c in leer_csv(raiz / CANDIDATOS, recortar=False):
        automatico = es_herramienta(c)
        filas.append({
            "id_fuente": c["id_fuente"],
            "corpus_actual": c["sugerencia_corpus"],
            "corpus_final": "descartados" if automatico else "revisar",
            "decision_revision": "descartar" if automatico else "revisar",
            "motivo_revision": "herramienta_textil" if automatico else "",
            "observaciones_revision": "Descarte automatico: herramienta de hilado o tejido." if automatico else "",
            "revisor": "auto" if automatico else "",
            "fecha_revision": c["fecha_descarga"] if automatico else "",
            "titulo_original": c["titulo_original"],
            "titulo_es_sugerido": "",
            "cultura": c["cultura"],
            "fecha_objeto": c["fecha_objeto"],
            "procedencia": c["procedencia"],
            "tipo_objeto": c["tipo_objeto"],
            "tipo_superficie": c["tipo_superficie"],
            "material": c["material"],
            "material_normalizado": "",
            "tecnica": c["tecnicas"],
            "estado_curacion": c["estado_filtro"],
            "destino_curacion": c["sugerencia_corpus"],
            "motivo_curacion": c["motivo_sugerencia"],
            "url_objeto": c["url_objeto"],
            "url_imagen": c["url_imagen_miniatura"] or c["url_imagen"],
            "ruta_imagen_local": "",
            "licencia": c["licencia_imagen"],
            "origen_curacion": "descarte_automatico_aic" if automatico else "revision_manual_aic",
        })
    return filas


def crear_hoja_resumen(libro: Workbook, filas: list[dict[str, str]]) -> None:
    hoja = libro.create_sheet("resumen", 0)
    hoja["A1"] = "Revision manual AIC v2"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A3"], hoja["B3"] = "Sugerencia automatica (corpus_actual)", "Registros"
    hoja["A3"].font = hoja["B3"].font = Font(bold=True)
    for fila_excel, (corpus, n) in enumerate(sorted(Counter(f["corpus_actual"] for f in filas).items()), start=4):
        hoja[f"A{fila_excel}"], hoja[f"B{fila_excel}"] = corpus, n
    instrucciones = [
        "Uso",
        "corpus_final: principal, secundario, descartados o revisar. Mientras diga revisar, el registro queda pendiente.",
        "corpus_actual es la sugerencia automatica (filtro textil y tipo de superficie); no se modifica.",
        "decision_revision, motivo_revision, revisor y fecha_revision documentan la decision humana.",
        "Las herramientas textiles (revisor = auto) se descartaron automaticamente, como en CMA.",
    ]
    for desplazamiento, texto in enumerate(instrucciones):
        hoja[f"A{10 + desplazamiento}"] = texto
    hoja["A10"].font = Font(bold=True)
    hoja.column_dimensions["A"].width = 110
    hoja.column_dimensions["B"].width = 14


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Crea el workbook maestro de revision manual AIC v2.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    parser.add_argument(
        "--forzar",
        action="store_true",
        help="Sobrescribe el workbook existente y pierde las decisiones manuales registradas.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    raiz = Path(args.root).resolve()
    salida = raiz / WORKBOOK
    if salida.exists() and not args.forzar:
        print(f"El workbook ya existe y contiene decisiones manuales: {salida}")
        print("No se sobrescribio. Use --forzar solo si quiere regenerarlo desde cero.")
        return 1

    filas = construir_filas(raiz)
    libro = Workbook()
    libro.remove(libro.active)
    crear_hoja_resumen(libro, filas)
    hoja = libro.create_sheet(HOJA_REVISION)
    hoja.append(COLUMNAS)
    for fila in filas:
        hoja.append([fila.get(columna, "") for columna in COLUMNAS])
    aplicar_estilo_revision(hoja)
    agregar_validaciones(hoja)
    agregar_hipervinculos(hoja)
    crear_hoja_listas(libro)

    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    print(f"Workbook creado: {salida}")
    print(f"Registros para revision: {len(filas)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
