"""Lectura de workbooks de Excel (openpyxl) con una sola convención para todo el proyecto.

Los workbooks de revisión y de anotación son el registro de las decisiones humanas. Se
leen siempre igual: cada fila como diccionario por encabezado, cada celda como texto
recortado, las celdas vacías como "" y las fechas como AAAA-MM-DD. Así un mismo workbook
produce los mismos valores en la aplicación de la revisión, la calibración, la
validación y el traslado de decisiones.
"""

from __future__ import annotations

from collections.abc import Callable
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def celda(valor: object) -> str:
    """Valor de una celda como texto: vacío -> "", fechas -> AAAA-MM-DD, el resto recortado."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor).strip()


def leer_hoja(ruta: Path, hoja: str, clave: str | None = None) -> list[dict[str, str]]:
    """Filas de una hoja como diccionarios por encabezado.

    clave: si se indica, se omiten las filas sin valor en esa columna (filas en blanco o
    de formato al final de la hoja).
    """
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ValueError(f"{ruta} no tiene la hoja requerida: {hoja}")
        filas = libro[hoja].iter_rows(values_only=True)
        encabezado = [celda(valor) for valor in next(filas)]
        salida = []
        for valores in filas:
            fila = {nombre: celda(valor) for nombre, valor in zip(encabezado, valores)}
            if clave is None or fila.get(clave):
                salida.append(fila)
        return salida
    finally:
        libro.close()


def lector(fila: dict[str, str]) -> Callable[[str], str]:
    """Función leer(columna) sobre una fila leída, para las reglas que reciben un lector."""
    return lambda columna: fila.get(columna, "")


def indice_columnas(hoja) -> dict[str, int]:
    """Número de columna (base 1) de cada encabezado de una hoja abierta para escritura."""
    return {celda(c.value): c.column for c in hoja[1] if c.value is not None}
