"""Traslado de decisiones del revisor desde un workbook de trabajo a los workbooks MET y CMA.

Los workbooks de revision (data/metadata/*_revision_manual_v2.xlsx) son el registro unico
de las decisiones curatoriales. Cuando el revisor decide en un workbook auxiliar (la
muestra de validacion ciega o la revision con el criterio v2), sus decisiones se trasladan
aqui con su motivo, su fecha y una observacion que cita el archivo de origen y la decision
anterior, de modo que ninguna se pierda ni se edite a mano en los CSV.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from src.imagenes.aptitud_workbooks import decision_para, motivo_vigente
from src.imagenes.motivos import analizar_motivo
from src.utils.fuentes import FUENTES, SUBCONJUNTOS
from src.utils.libros import celda, indice_columnas, leer_hoja

REVISOR = "Henry Sánchez"


def contradice_motivo(decision: str, motivo: str) -> bool:
    """La decision contradice el criterio del 2026-10-19 segun el propio motivo del revisor.

    Principal = colores y formas identificables (completo o incompleto). Un descarte o un
    secundario cuyo motivo dice «con colores y formas visibles», o un principal cuyo motivo
    dice que no se ven, se deja en revisar para que el revisor lo resuelva: el sistema no
    elige entre dos cosas que el revisor afirmo a la vez.
    """
    a = analizar_motivo(motivo)
    legible = a["colores"] == "si" and a["formas"] == "si"
    ilegible = a["colores"] == "no" or a["formas"] == "no"
    return (decision != "principal" and legible and not a["pertinencia"]) or (decision == "principal" and ilegible)


def trasladar_anotaciones(raiz: Path, archivo: str, hoja_anotaciones: str, fecha: str, origen: str, descripcion: str,
                          registrar_confirmaciones: bool) -> Counter:
    """Lleva las decisiones anotadas en `archivo` a los workbooks MET y CMA.

    - Una anotacion vacia o fuera de principal/secundario/descartados no cambia nada.
    - Si contradice su propio motivo, el registro queda en revisar con la explicacion.
    - registrar_confirmaciones: si la decision coincide con la vigente, igual se registra el
      motivo nuevo (la revision v2 confirma decisiones con el criterio vigente). La validacion
      ciega no lo hace: alli solo importan los cambios.
    - Es idempotente: una fila ya trasladada con la misma decision y el mismo motivo se omite.
    """
    anotaciones = {fila["id_global"]: fila for fila in leer_hoja(raiz / archivo, hoja_anotaciones, clave="id_global")}
    conteo: Counter = Counter()
    for fuente in ("MET", "CMA"):
        ruta = raiz / FUENTES[fuente].workbook
        libro = load_workbook(ruta)
        hoja = libro[FUENTES[fuente].hoja]
        indice = indice_columnas(hoja)
        cambios = 0
        for fila in range(2, hoja.max_row + 1):
            id_global = f"{fuente}:{celda(hoja.cell(row=fila, column=indice['id_fuente']).value)}"
            anotacion = anotaciones.get(id_global)
            if not anotacion or anotacion.get("corpus_final") not in SUBCONJUNTOS:
                continue

            def leer(columna: str) -> str:
                return celda(hoja.cell(row=fila, column=indice[columna]).value)

            decision, motivo = anotacion["corpus_final"], anotacion.get("motivo_revision", "")
            conflicto = contradice_motivo(decision, motivo)
            destino = "revisar" if conflicto else decision
            if leer("origen_curacion") == origen and (not registrar_confirmaciones or
                                                      (leer("corpus_final") == destino and leer("motivo_revision") == motivo)):
                continue  # ya trasladada
            confirma = destino == leer("corpus_final")
            if confirma and not (registrar_confirmaciones and motivo):
                conteo["coinciden"] += 1
                continue
            anterior = motivo_vigente(leer) or "sin motivo"
            nota = (f"{descripcion} del {fecha} ({archivo}, orden {anotacion.get('orden', '')}): {decision}. "
                    f"Decisión anterior: {leer('corpus_final')} ({anterior}).")
            if conflicto:
                nota += " Queda en revisar: la anotación contradice su propio motivo según el criterio del 2026-10-19."
            valores = {
                "corpus_final": destino,
                "decision_revision": "revisar" if conflicto else decision_para(leer("corpus_actual"), destino),
                "motivo_revision": motivo,
                "observaciones_revision": nota,
                "revisor": REVISOR,
                "fecha_revision": fecha,
                "origen_curacion": origen,
            }
            for columna, valor in valores.items():
                hoja.cell(row=fila, column=indice[columna], value=valor)
            clase = "revisar" if conflicto else ("confirmadas" if confirma else "cambiadas")
            conteo[clase] += 1
            cambios += 1
            print(f"  {id_global}: {valores['corpus_final']} ({clase})")
        if cambios:
            libro.save(ruta)  # guardar sin cambios reescribe el xlsx y ensucia el historial de git
    return conteo
