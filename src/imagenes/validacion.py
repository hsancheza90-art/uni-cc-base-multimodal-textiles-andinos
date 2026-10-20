"""Muestra de validacion ciega del puntaje de aptitud en MET y CMA.

Uso desde la raiz del repositorio (sin red):

    python -m src.imagenes.validacion --crear      # muestra y workbook ciego (no sobrescribe sin --forzar)
    python -m src.imagenes.validacion              # evalua las anotaciones del workbook
    python -m src.imagenes.validacion --trasladar  # lleva las anotaciones a los workbooks MET y CMA

El modelo se calibro primero solo con las decisiones AIC. Antes de aplicarlo a MET y CMA se
comprueba si transfiere: el revisor anota una muestra con el criterio AIC
(principal, secundario o descartados y su motivo) sin ver el puntaje ni la
decision v1, y se mide el acuerdo con la banda sugerida y con la decision v1.

El informe versionado (outputs/reports/validacion_met_cma_v2.md) registra esa medicion
del 2026-10-19, hecha con el modelo calibrado solo con AIC. Tras el traslado, las
decisiones v1 de la muestra ya son las del revisor y la comparacion pierde sentido: por
eso la evaluacion no se vuelve a escribir una vez trasladadas las anotaciones.

La muestra excluye lo que ya deciden reglas verificables (pertinencia por
metadatos, fotografia monocroma, resolucion, sin imagen) y los descartes v1 por
pertinencia (no andino, no textil, herramienta, materia prima, duplicado): ahi el
modelo de imagen no interviene. Se estratifica por fuente y banda sugerida, con
semilla fija; el orden de la galeria y del workbook es aleatorio.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.imagenes.aptitud import SALIDA_APTITUD
from src.imagenes.aptitud_workbooks import PUERTAS, es_descarte_pertinencia_v1
from src.imagenes.descargar_evaluacion import MANIFIESTO_EVALUACION
from src.imagenes.motivos import analizar_motivo
from src.imagenes.traslado import trasladar_anotaciones
from src.review.galerias import esc, escribir_pagina, imagen_evaluacion, tarjeta
from src.utils.archivos import escribir_texto, indexar, leer_csv
from src.utils.fuentes import FUENTES, SUBCONJUNTOS
from src.utils.libros import leer_hoja, lector
from src.utils.metricas import kappa, tabla

WORKBOOK_VALIDACION = "data/metadata/validacion_met_cma_v2.xlsx"
GALERIA_VALIDACION = "outputs/review/validacion_met_cma_v2_galeria.html"
INFORME_VALIDACION = "outputs/reports/validacion_met_cma_v2.md"
HOJA = "validacion"
SEMILLA = 20261019
TAMANO_POR_FUENTE = {"MET": 20, "CMA": 20}
MINIMO_POR_BANDA = 4
BANDAS = ("principal", "revisar", "descartar")

COLUMNAS = ["orden", "id_global", "titulo_original", "cultura", "url_objeto", "url_imagen",
            "corpus_final", "motivo_revision", "revisor", "fecha_revision"]


def decisiones_v1(raiz: Path, fuente: str) -> dict[str, dict[str, str]]:
    salida = {}
    for valor in leer_hoja(raiz / FUENTES[fuente].workbook, FUENTES[fuente].hoja, clave="id_fuente"):
        valor["pertinencia_v1"] = "si" if es_descarte_pertinencia_v1(lector(valor)) else "no"
        salida[f"{fuente}:{valor['id_fuente']}"] = valor
    return salida


def universo(raiz: Path) -> list[dict[str, str]]:
    """Registros MET y CMA donde decide el modelo de imagen (sin puertas ni pertinencia v1)."""
    aptitud = indexar(leer_csv(raiz / SALIDA_APTITUD), "id_global")
    salida = []
    for fuente in TAMANO_POR_FUENTE:
        for id_global, v1 in decisiones_v1(raiz, fuente).items():
            ficha = aptitud.get(id_global)
            if not ficha or ficha["motivo_banda"] in PUERTAS or v1["pertinencia_v1"] == "si":
                continue
            if v1["corpus_final"] not in SUBCONJUNTOS:
                continue
            salida.append({"id_global": id_global, "fuente": fuente, "banda": ficha["banda_sugerida"],
                           "decision_v1": v1["corpus_final"], "titulo_original": v1["titulo_original"],
                           "cultura": v1["cultura"], "url_objeto": v1["url_objeto"], "url_imagen": v1["url_imagen"]})
    return salida


def muestrear(candidatos: list[dict[str, str]]) -> list[dict[str, str]]:
    """Muestra estratificada por fuente y banda: minimo por banda, el resto proporcional."""
    rng = np.random.default_rng(SEMILLA)
    muestra = []
    for fuente, tamano in TAMANO_POR_FUENTE.items():
        estratos = {b: sorted((c for c in candidatos if c["fuente"] == fuente and c["banda"] == b), key=lambda c: c["id_global"])
                    for b in BANDAS}
        cupos = {b: min(len(v), MINIMO_POR_BANDA) for b, v in estratos.items()}
        restante = tamano - sum(cupos.values())
        total = sum(len(v) - cupos[b] for b, v in estratos.items())
        for b, v in estratos.items():
            if total:
                cupos[b] += min(len(v) - cupos[b], round(restante * (len(v) - cupos[b]) / total))
        # El redondeo puede dejar la fuente corta o larga: el ajuste va al estrato con mas disponibles.
        diferencia = tamano - sum(cupos.values())
        mayor = max(estratos, key=lambda b: len(estratos[b]) - cupos[b])
        cupos[mayor] = max(0, min(len(estratos[mayor]), cupos[mayor] + diferencia))
        for b, v in estratos.items():
            indices = rng.choice(len(v), cupos[b], replace=False) if cupos[b] else []
            muestra += [v[i] for i in sorted(indices)]
    orden = rng.permutation(len(muestra))
    return [dict(muestra[i], orden=str(n + 1)) for n, i in enumerate(orden)]


def crear(raiz: Path, forzar: bool) -> int:
    salida = raiz / WORKBOOK_VALIDACION
    if salida.exists() and not forzar:
        print(f"{WORKBOOK_VALIDACION} ya existe y puede tener anotaciones; no se sobrescribio (use --forzar).")
        return 1
    muestra = muestrear(universo(raiz))
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA
    hoja.append(COLUMNAS)
    for fila in muestra:
        hoja.append([fila.get(c, "") for c in COLUMNAS])
    for celda in hoja[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    for letra, ancho in zip("ABCDEFGHIJ", (7, 13, 36, 22, 40, 40, 14, 60, 14, 12)):
        hoja.column_dimensions[letra].width = ancho
    for fila in range(2, hoja.max_row + 1):
        for letra in "EF":
            if hoja[f"{letra}{fila}"].value:
                hoja[f"{letra}{fila}"].hyperlink = hoja[f"{letra}{fila}"].value
    validacion = DataValidation(type="list", formula1='"principal,secundario,descartados"', allow_blank=True)
    hoja.add_data_validation(validacion)
    validacion.add(f"G2:G{hoja.max_row}")
    hoja.freeze_panes = "B2"
    instrucciones = libro.create_sheet("instrucciones")
    for i, texto in enumerate([
        "Validacion ciega del puntaje de aptitud en MET y CMA",
        "Anote cada imagen con el mismo criterio que uso en AIC: principal si los colores y las formas se identifican "
        "(aunque la pieza este incompleta), secundario si tiene muchos defectos, descartados si no sirve para el estudio.",
        "Escriba el motivo como en AIC (por ejemplo: Textil incompleto con colores y formas visibles).",
        "La muestra no muestra el puntaje ni la decision anterior a proposito: no los busque hasta terminar.",
        f"Galeria con las mismas imagenes y el mismo orden: {GALERIA_VALIDACION}",
    ], start=1):
        instrucciones[f"A{i}"] = texto
    instrucciones["A1"].font = Font(bold=True, size=14)
    instrucciones.column_dimensions["A"].width = 120
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    galeria(raiz, muestra)
    print(f"Muestra: {len(muestra)} registros -> {salida}")
    print("Por fuente y banda (no se muestra al revisor):", dict(Counter((m["fuente"], m["banda"]) for m in muestra)))
    return 0


def galeria(raiz: Path, muestra: list[dict[str, str]]) -> None:
    """Galeria ciega: solo imagen, titulo y cultura (sin puntaje ni decision anterior)."""
    salida = raiz / GALERIA_VALIDACION
    manifiesto = indexar(leer_csv(raiz / MANIFIESTO_EVALUACION), "id_global")
    tarjetas = [
        tarjeta(imagen_evaluacion(raiz, manifiesto, f["id_global"], salida), f["url_objeto"], f["id_global"], [
            f"<b>{esc(f['orden'])}.</b> {esc(f['id_global'])}",
            ("t", esc(f["titulo_original"])),
            esc(f["cultura"]),
        ])
        for f in sorted(muestra, key=lambda f: int(f["orden"]))
    ]
    nota = (f"Anote cada imagen en <code>{WORKBOOK_VALIDACION}</code> con el criterio AIC. "
            "No se muestran el puntaje ni la decisión anterior.")
    escribir_pagina(salida, f"Validación ciega: {len(muestra)} imágenes MET y CMA", nota, [("", tarjetas)])


# ------------------------------------------------------------------ evaluacion

def evaluar(raiz: Path) -> int:
    filas = leer_hoja(raiz / WORKBOOK_VALIDACION, HOJA, clave="id_global")
    anotadas = [a for a in filas if a["corpus_final"] in SUBCONJUNTOS]
    aptitud = indexar(leer_csv(raiz / SALIDA_APTITUD), "id_global")
    v1 = {**decisiones_v1(raiz, "MET"), **decisiones_v1(raiz, "CMA")}
    if any(fila.get("origen_curacion") == ORIGEN_VALIDACION for fila in v1.values()):
        print(f"Las anotaciones ya se trasladaron a los workbooks: {INFORME_VALIDACION} se conserva como registro de la validacion.")
        return 0
    if not anotadas:
        print("La muestra aun no tiene anotaciones.")
        return 0

    def principal(valor: str) -> str:
        return "principal" if valor == "principal" else "no_principal"

    lineas = ["# Validación ciega del puntaje de aptitud en MET y CMA", "",
              f"Generado por `python -m src.imagenes.validacion` desde `{WORKBOOK_VALIDACION}` "
              f"({len(anotadas)} de {len(filas)} imágenes anotadas).", "",
              "El revisor anotó sin ver el puntaje ni la decisión v1, con el criterio aplicado en AIC. "
              "La comparación principal es binaria (principal frente al resto), porque el puntaje no separa "
              "secundario de descartados y la zona «revisar» queda para la decisión manual.", ""]
    filas_tabla = []
    for nombre, filtro in (("MET + CMA", ("MET", "CMA")), ("MET", ("MET",)), ("CMA", ("CMA",))):
        grupo = [a for a in anotadas if a["id_global"].split(":")[0] in filtro]
        con_banda = [a for a in grupo if aptitud[a["id_global"]]["banda_sugerida"] != "revisar"]
        pares_modelo = [(principal(a["corpus_final"]), principal(aptitud[a["id_global"]]["banda_sugerida"])) for a in con_banda]
        pares_v1 = [(principal(a["corpus_final"]), principal(v1[a["id_global"]]["corpus_final"])) for a in grupo]
        acuerdo = sum(x == y for x, y in pares_modelo) / len(pares_modelo) if pares_modelo else float("nan")
        filas_tabla.append([nombre, str(len(grupo)), str(len(con_banda)), f"{acuerdo:.2f}", f"{kappa(pares_modelo):.2f}",
                            f"{sum(x == y for x, y in pares_v1) / len(pares_v1):.2f}" if pares_v1 else "", f"{kappa(pares_v1):.2f}"])
    lineas += tabla(filas_tabla, ["Conjunto", "Anotadas", "Con banda (sin revisar)", "Acuerdo con el modelo",
                                  "Kappa con el modelo", "Acuerdo con v1", "Kappa con v1"])
    cruce = Counter((a["corpus_final"], aptitud[a["id_global"]]["banda_sugerida"]) for a in anotadas)
    lineas += ["", "## Anotación frente a banda sugerida", ""]
    lineas += tabla([[d] + [str(cruce[(d, b)]) for b in BANDAS] for d in SUBCONJUNTOS],
                    ["Anotación \\ banda", *BANDAS])
    revisar = [a for a in anotadas if aptitud[a["id_global"]]["banda_sugerida"] == "revisar"]
    lineas += ["", f"En la banda «revisar» ({len(revisar)} imágenes) el revisor puso: "
               + ", ".join(f"{d} {n}" for d, n in Counter(a["corpus_final"] for a in revisar).most_common()) + "."]
    desacuerdos = [a for a in anotadas if aptitud[a["id_global"]]["banda_sugerida"] != "revisar"
                   and principal(a["corpus_final"]) != principal(aptitud[a["id_global"]]["banda_sugerida"])]
    lineas += ["", "## Desacuerdos con el modelo", ""]
    lineas += tabla([[a["id_global"], a["corpus_final"], aptitud[a["id_global"]]["banda_sugerida"],
                      aptitud[a["id_global"]]["puntaje_aptitud"], a["motivo_revision"][:60]] for a in desacuerdos],
                    ["Registro", "Anotación", "Banda", "Puntaje", "Motivo del revisor"]) if desacuerdos else ["Ninguno."]
    atributos = Counter()
    for a in anotadas:
        m = analizar_motivo(a["motivo_revision"])
        for clave in ("colores", "formas", "completitud"):
            if m[clave]:
                atributos[(clave, m[clave], aptitud[a["id_global"]][clave])] += 1
    lineas += ["", "## Atributos: anotación del revisor frente a nivel estimado", ""]
    lineas += tabla([[c, v, n_, str(n)] for (c, v, n_), n in sorted(atributos.items())],
                    ["Atributo", "Revisor", "Imagen", "Registros"])
    ruta = raiz / INFORME_VALIDACION
    escribir_texto(ruta, lineas)
    print(f"Informe: {ruta}")
    return 0


# ------------------------------------------------------------------ traslado a los workbooks

ORIGEN_VALIDACION = "validacion_ciega_v2"


def trasladar(raiz: Path, fecha: str) -> int:
    """Lleva las anotaciones de la validacion a los workbooks MET y CMA como decisiones del revisor."""
    conteo = trasladar_anotaciones(raiz, WORKBOOK_VALIDACION, HOJA, fecha, ORIGEN_VALIDACION,
                                   "Anotación de la validación ciega", registrar_confirmaciones=False)
    print(f"Cambiadas {conteo['cambiadas']} | en revisar por contradicción {conteo['revisar']} | ya coincidían {conteo['coinciden']}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Muestra de validacion ciega del puntaje de aptitud en MET y CMA.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    parser.add_argument("--crear", action="store_true", help="Crea la muestra, el workbook ciego y la galeria.")
    parser.add_argument("--forzar", action="store_true", help="Con --crear, sobrescribe el workbook (pierde anotaciones).")
    parser.add_argument("--trasladar", action="store_true", help="Lleva las anotaciones a los workbooks MET y CMA.")
    parser.add_argument("--fecha", default="2026-10-19", help="Fecha de las decisiones trasladadas.")
    args = parser.parse_args(argv)
    raiz = Path(args.root).resolve()
    if args.trasladar:
        return trasladar(raiz, args.fecha)
    return crear(raiz, args.forzar) if args.crear else evaluar(raiz)


if __name__ == "__main__":
    raise SystemExit(main())
