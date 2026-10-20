"""Revision de MET y CMA con el criterio v2 de aptitud analitica.

Uso desde la raiz del repositorio (sin red):

    python -m src.imagenes.revision_v2 --crear       # workbook y galeria (no sobrescribe sin --forzar)
    python -m src.imagenes.revision_v2 --trasladar   # lleva las decisiones a los workbooks MET y CMA

Parte de MET y CMA se decidio con criterios anteriores: el secundario por tipo de objeto
(prendas tridimensionales, fajas, sombreros) y los descartes por «fragmento no legible» o
«baja calidad». El criterio v2 (docs/protocolo_aptitud_v2.md) puede cambiar esas
decisiones: principal si los colores y las figuras se identifican a primera vista, aunque
la pieza este incompleta. Este modulo reune esos registros en un solo workbook para que el
revisor los decida, con la ficha de aptitud como apoyo (la validacion ciega mostro que la
banda principal acierta el 91 % en MET y CMA, pero que el extremo bajo es mas estricto que
el revisor: la sugerencia orienta, no decide).

Se incluyen: secundarios, descartes no debidos a pertinencia, pendientes, y principales
cuya ficha sugiere descartar. Se excluyen: decisiones automaticas por reglas, las ya
tomadas con el criterio v2 en la validacion ciega y los descartes v1 por pertinencia (no
andino, no textil, herramienta, duplicado), que no dependen de la imagen.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.imagenes.aptitud import SALIDA_APTITUD
from src.imagenes.aptitud_workbooks import es_decision_automatica, es_descarte_pertinencia_v1, motivo_vigente
from src.imagenes.descargar_evaluacion import MANIFIESTO_EVALUACION
from src.imagenes.traslado import trasladar_anotaciones
from src.imagenes.validacion import ORIGEN_VALIDACION
from src.review.galerias import esc, escribir_pagina, imagen_evaluacion, tarjeta
from src.utils.archivos import indexar, leer_csv
from src.utils.fuentes import FUENTES
from src.utils.libros import leer_hoja, lector

WORKBOOK_REVISION = "data/metadata/revision_criterio_v2_met_cma.xlsx"
GALERIA_REVISION = "outputs/review/revision_criterio_v2_met_cma_galeria.html"
HOJA = "revision"
ORIGEN_REVISION = "revision_criterio_v2"
ORDEN_DECISION = {"secundario": 0, "descartados": 1, "revisar": 2, "principal": 3}

# Columnas que decide el revisor, a la izquierda; la ficha de apoyo, a la derecha.
COLUMNAS = [
    ("orden", 7), ("id_global", 13), ("titulo_original", 34), ("cultura", 18), ("decision_actual", 14),
    ("corpus_final", 14), ("motivo_revision", 50),
    ("motivo_actual", 34), ("banda_sugerida", 13), ("referencia_sugerida", 12), ("puntaje_aptitud", 10),
    ("fotografia", 12), ("colores", 16), ("formas", 14), ("completitud", 13), ("justificacion_aptitud", 80),
    ("url_objeto", 40), ("url_imagen", 40),
]
COLUMNAS_DECISION = {"corpus_final", "motivo_revision"}


def seleccionar(raiz: Path) -> list[dict[str, str]]:
    """Registros MET y CMA que conviene decidir de nuevo con el criterio v2."""
    aptitud = indexar(leer_csv(raiz / SALIDA_APTITUD), "id_global")
    salida = []
    for fuente in ("MET", "CMA"):
        for fila in leer_hoja(raiz / FUENTES[fuente].workbook, FUENTES[fuente].hoja, clave="id_fuente"):
            leer = lector(fila)
            id_global = f"{fuente}:{leer('id_fuente')}"
            ficha = aptitud.get(id_global, {})
            decision = leer("corpus_final")
            ya_decidida_v2 = leer("origen_curacion") in (ORIGEN_VALIDACION, ORIGEN_REVISION)
            if es_decision_automatica(leer) or ya_decidida_v2 or es_descarte_pertinencia_v1(leer):
                continue  # ya decididas con el criterio v2, o no dependen de la imagen
            if decision == "principal" and ficha.get("banda_sugerida") != "descartar":
                continue  # el criterio v2 es mas permisivo: un principal solo se revisa si la imagen lo pone en duda
            motivo = motivo_vigente(leer)
            salida.append({
                "id_global": id_global, "fuente": fuente, "titulo_original": leer("titulo_original"), "cultura": leer("cultura"),
                "decision_actual": decision, "motivo_actual": motivo, "url_objeto": leer("url_objeto"), "url_imagen": leer("url_imagen"),
                **{c: ficha.get(c, "") for c in ("banda_sugerida", "referencia_sugerida", "puntaje_aptitud", "fotografia",
                                                  "colores", "formas", "completitud")},
                "justificacion_aptitud": ficha.get("justificacion", ""),
            })
    salida.sort(key=lambda f: (f["fuente"] != "MET", ORDEN_DECISION.get(f["decision_actual"], 9), -float(f["puntaje_aptitud"] or 0)))
    for numero, fila in enumerate(salida, start=1):
        fila["orden"] = str(numero)
    return salida


def crear(raiz: Path, forzar: bool) -> int:
    salida = raiz / WORKBOOK_REVISION
    if salida.exists() and not forzar:
        print(f"{WORKBOOK_REVISION} ya existe y puede tener decisiones; no se sobrescribio (use --forzar).")
        return 1
    registros = seleccionar(raiz)
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA
    hoja.append([c for c, _ in COLUMNAS])
    for registro in registros:
        hoja.append([float(registro[c]) if c == "puntaje_aptitud" and registro[c] else registro.get(c, "") for c, _ in COLUMNAS])
    for numero, (columna, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=numero)
        celda.font = Font(color="FFFFFF", bold=True)
        # Verde: decide el revisor; azul: datos del registro; ocre: ficha de apoyo.
        color = "38761D" if columna in COLUMNAS_DECISION else ("7F6000" if numero > 7 and columna not in ("url_objeto", "url_imagen") else "1F4E78")
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[celda.column_letter].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        for celda in fila[-2:]:
            if celda.value:
                celda.hyperlink, celda.style = celda.value, "Hyperlink"
    validacion = DataValidation(type="list", formula1='"principal,secundario,descartados"', allow_blank=True)
    hoja.add_data_validation(validacion)
    validacion.add(f"F2:F{hoja.max_row}")
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions

    instrucciones = libro.create_sheet("instrucciones")
    textos = [
        "Revisión de MET y CMA con el criterio v2",
        "Principal: los colores y las figuras se identifican a primera vista, esté completo o no el textil.",
        "Imagen de referencia: además está completo y se verifica en la fotografía (escriba «completo» en el motivo).",
        "Secundario: la imagen no permite identificar bien colores o formas. Descarte: no sirve o no es pertinente.",
        "Complete corpus_final (lista) y motivo_revision, en verde. Si deja corpus_final vacío, la decisión actual no cambia.",
        "Escriba el motivo como en AIC, por ejemplo: «Textil incompleto con colores y formas visibles».",
        "Las columnas ocres son la ficha de aptitud: orienta, no decide. Galería con el mismo orden: " + GALERIA_REVISION,
        "Al terminar: python -m src.imagenes.revision_v2 --trasladar y luego python -m src.flujo.",
    ]
    for numero, texto in enumerate(textos, start=1):
        instrucciones[f"A{numero}"] = texto
    instrucciones["A1"].font = Font(bold=True, size=14)
    instrucciones.column_dimensions["A"].width = 130
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    galeria(raiz, registros)
    print(f"Revisión: {len(registros)} registros -> {salida}")
    print("Por fuente y decisión actual:", dict(Counter((r["fuente"], r["decision_actual"]) for r in registros)))
    return 0


def galeria(raiz: Path, registros: list[dict[str, str]]) -> None:
    """Galeria local (no se versiona) con la imagen de evaluacion y la ficha, en el orden del workbook."""
    salida = raiz / GALERIA_REVISION
    manifiesto = indexar(leer_csv(raiz / MANIFIESTO_EVALUACION), "id_global")
    tarjetas = []
    for r in registros:
        referencia = " · referencia" if r["referencia_sugerida"] == "si" else ""
        tarjetas.append(tarjeta(imagen_evaluacion(raiz, manifiesto, r["id_global"], salida), r["url_objeto"], r["id_global"], [
            f"<b>{esc(r['orden'])}.</b> {esc(r['id_global'])} <span class=\"p\">{esc(r['puntaje_aptitud'])}</span>",
            ("t", f"{esc(r['titulo_original'])} · {esc(r['cultura'])}"),
            f"<b>Actual:</b> {esc(r['decision_actual'])} <span class=\"m\">({esc(r['motivo_actual'])})</span>",
            f"<b>Sugerencia:</b> {esc(r['banda_sugerida'])}{referencia} · "
            + " · ".join(esc(r[c]) for c in ("fotografia", "colores", "formas", "completitud")),
            ("m", esc(r["justificacion_aptitud"])),
        ]))
    nota = f"Mismo orden que <code>{WORKBOOK_REVISION}</code>. Decida en el workbook; la sugerencia orienta, no decide."
    escribir_pagina(salida, f"Revisión de MET y CMA con el criterio v2 ({len(registros)} registros)", nota, [("", tarjetas)])


def trasladar(raiz: Path, fecha: str) -> int:
    conteo = trasladar_anotaciones(raiz, WORKBOOK_REVISION, HOJA, fecha, ORIGEN_REVISION,
                                   "Revisión con el criterio v2", registrar_confirmaciones=True)
    print(f"Cambiadas {conteo['cambiadas']} | confirmadas con motivo nuevo {conteo['confirmadas']} | "
          f"en revisar por contradicción {conteo['revisar']}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Revisión de MET y CMA con el criterio v2 de aptitud analítica.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    parser.add_argument("--crear", action="store_true", help="Crea el workbook de revisión y la galería.")
    parser.add_argument("--forzar", action="store_true", help="Con --crear, sobrescribe el workbook (pierde las decisiones).")
    parser.add_argument("--trasladar", action="store_true", help="Lleva las decisiones a los workbooks MET y CMA.")
    parser.add_argument("--fecha", default=date.today().isoformat(), help="Fecha de las decisiones trasladadas.")
    args = parser.parse_args(argv)
    raiz = Path(args.root).resolve()
    if args.crear:
        return crear(raiz, args.forzar)
    if args.trasladar:
        return trasladar(raiz, args.fecha)
    parser.print_help()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
