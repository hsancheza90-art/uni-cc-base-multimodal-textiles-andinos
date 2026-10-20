"""Calibra la ficha de aptitud con las decisiones del revisor en AIC, MET y CMA.

Uso desde la raiz del repositorio (sin red; tambien lo ejecuta python -m src.flujo):

    python -m src.imagenes.calibrar_aptitud

Lee las medidas de imagen (data/metadata/aptitud_medidas_v2.csv) y las decisiones de
los workbooks de revision. Del motivo libre del revisor extrae los atributos
(src/imagenes/motivos.py) y ajusta:

1. Un modelo por atributo (regresion logistica con L2) que estima, desde la imagen,
   si los colores son identificables, si las formas son visibles, si el textil esta
   completo y si es una imagen de referencia para el estudio.
2. Un puntaje ordinal (0-100): media ponderada, con pesos no negativos, de la calidad
   fotografica y de las probabilidades de los atributos, ajustada para ordenar
   referencia > principal > secundario > descarte por imagen.

Que decisiones calibran ([calibracion] de config/aptitud.toml): todas las humanas de
AIC, y de MET y CMA solo las tomadas con el criterio v2 (validacion ciega y revision
v2); el resto de MET y CMA se decidio con criterios anteriores. Calibrar con las tres
fuentes es necesario porque el modelo calibrado solo con AIC no generalizaba a MET
(separacion util/resto 0,49 frente a 0,76 con la calibracion conjunta).

Las etiquetas del revisor no son definitivas (hay piezas completas sin anotar como
tales): por eso los modelos usan regularizacion, y la validacion cruzada se estratifica
por fuente y clase para que cada pliegue represente a las tres colecciones.

Escribe el modelo en config/aptitud_modelo.json y el informe en
outputs/reports/aptitud_calibracion_v2.md.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from src.imagenes.aptitud import (
    PERTINENCIA_DESCARTE,
    capa_calidad,
    metadatos_pertinencia,
    motivo_pertinencia,
    transformar,
    valor_medida,
    vector,
)
from src.imagenes.aptitud_workbooks import WORKBOOKS, es_decision_automatica, motivo_del_revisor
from src.imagenes.medidas import MEDIDAS
from src.imagenes.motivos import analizar_motivo
from src.utils.archivos import indexar, leer_csv
from src.utils.config import leer_config_aptitud
from src.utils.metricas import auc

MODELO = "config/aptitud_modelo.json"
CLASES = ("descarte_imagen", "secundario", "principal", "referencia")
PLIEGUES = 5
SEMILLA = 0


# ------------------------------------------------------------------ etiquetas

def etiquetas(raiz: Path, fuente: str, medidas: dict[str, dict[str, str]], config: dict[str, Any],
              origenes: set[str] | None = None) -> list[dict[str, Any]]:
    """Decision y atributos del revisor para cada registro con imagen util y sin regla de pertinencia.

    origenes: si se indica, solo cuentan las filas cuyo origen_curacion esta en el conjunto
    (por ejemplo, las decididas con el criterio v2).
    """
    ruta, hoja = WORKBOOKS[fuente]
    libro = load_workbook(raiz / ruta, read_only=True, data_only=True)
    filas = list(libro[hoja].iter_rows(values_only=True))
    libro.close()
    indice = {nombre: i for i, nombre in enumerate(filas[0])}
    metadatos = metadatos_pertinencia(raiz)
    salida = []
    for fila in filas[1:]:
        def leer(columna: str, fila=fila) -> str:
            return str(fila[indice[columna]] or "").strip() if columna in indice else ""

        id_global = f"{fuente}:{leer('id_fuente')}"
        decision = leer("corpus_final")
        m = medidas.get(id_global)
        if decision not in ("principal", "secundario", "descartados") or not m or m["estado_imagen"] != "ok":
            continue
        if es_decision_automatica(leer):
            continue  # solo las decisiones humanas calibran
        if origenes is not None and leer("origen_curacion") not in origenes:
            continue
        atributos = analizar_motivo(motivo_del_revisor(leer))
        if atributos["pertinencia"] or atributos["monocroma"]:
            continue  # decisiones que no dependen de lo que la imagen permite ver
        if motivo_pertinencia(metadatos.get(id_global), config["pertinencia"]) in PERTINENCIA_DESCARTE:
            continue
        if float(m["croma_p90"]) < config["puertas"]["croma_monocromo"]:
            continue
        clase = {"principal": "principal", "secundario": "secundario", "descartados": "descarte_imagen"}[decision]
        if atributos["referencia"] and decision == "principal":
            # Criterio del 2026-10-19: principal y textil completo verificable en la fotografia.
            clase = "referencia"
        salida.append({"id_global": id_global, "fuente": fuente, "clase": clase, **atributos})
    return salida


# ------------------------------------------------------------------ modelos

def logistica(x: np.ndarray, y: np.ndarray, l2: float, iteraciones: int = 50) -> np.ndarray:
    """Regresion logistica con L2 (sin penalizar el intercepto), por Newton-Raphson."""
    xa = np.hstack([np.ones((len(x), 1)), x])
    w = np.zeros(xa.shape[1])
    penal = np.full(xa.shape[1], l2)
    penal[0] = 0
    for _ in range(iteraciones):
        p = 1 / (1 + np.exp(-np.clip(xa @ w, -30, 30)))
        grad = xa.T @ (p - y) + penal * w
        hess = xa.T @ (xa * (p * (1 - p))[:, None]) + np.diag(penal + 1e-9)
        paso = np.linalg.solve(hess, grad)
        w -= paso
        if np.abs(paso).max() < 1e-8:
            break
    return w


def predecir(x: np.ndarray, w: np.ndarray) -> np.ndarray:
    return 1 / (1 + np.exp(-np.clip(w[0] + x @ w[1:], -30, 30)))


def pliegues(clases: list[str]) -> np.ndarray:
    """Asignacion estratificada por clase a PLIEGUES pliegues, con semilla fija."""
    rng = np.random.default_rng(SEMILLA)
    salida = np.zeros(len(clases), dtype=int)
    for clase in sorted(set(clases)):
        indices = np.array([i for i, c in enumerate(clases) if c == clase])
        rng.shuffle(indices)
        salida[indices] = np.arange(len(indices)) % PLIEGUES
    return salida


def ordinal(z: np.ndarray, y: np.ndarray, iteraciones: int = 3000, paso: float = 0.5) -> tuple[np.ndarray, np.ndarray]:
    """Modelo de odds proporcionales con pesos no negativos: P(Y > k) = sigma(z.w - theta_k)."""
    k = int(y.max())
    w = np.full(z.shape[1], 1.0)
    theta = np.linspace(-1, 1, k)
    for _ in range(iteraciones):
        s = z @ w
        grad_w, grad_t = np.zeros_like(w), np.zeros_like(theta)
        for j in range(k):
            p = 1 / (1 + np.exp(-(s - theta[j])))
            objetivo = (y > j).astype(float)
            grad_w += z.T @ (p - objetivo) / len(y)
            grad_t -= np.sum(p - objetivo) / len(y)
        w = np.maximum(w - paso * grad_w, 0)
        theta = np.sort(theta - paso * grad_t)
    return w, theta


def medidas_atributo(config: dict[str, Any], atributo: str) -> list[str]:
    """Medidas propias del atributo si las declara; si no, las comunes de [modelo]."""
    return config["atributos"][atributo].get("medidas", config["modelo"]["medidas"])


class Ajuste:
    """Modelos por atributo y pesos del puntaje ajustados sobre un conjunto de registros."""

    def __init__(self, config: dict[str, Any], medidas: dict[str, dict[str, str]], registros: list[dict[str, Any]], estandar: dict):
        self.config, self.medidas, self.estandar = config, medidas, estandar
        self.atributos: dict[str, dict[str, Any]] = {}
        for nombre, definicion in config["atributos"].items():
            positivos, negativos = definicion["positivo"], definicion["negativo"]
            usados = [r for r in registros if r[definicion["etiqueta"]] in positivos + negativos]
            if not usados:
                continue
            nombres = medidas_atributo(config, nombre)
            x = np.array([vector(self.medidas[r["id_global"]], nombres, estandar) for r in usados])
            y = np.array([1.0 if r[definicion["etiqueta"]] in positivos else 0.0 for r in usados])
            pesos = [round(float(w), 6) for w in logistica(x, y, definicion["l2"])]
            self.atributos[nombre] = {"medidas": nombres, "pesos": pesos, "n": len(usados)}
        z = np.array([self.componentes(r["id_global"]) for r in registros])
        y = np.array([CLASES.index(r["clase"]) for r in registros])
        self.pesos, self.cortes = ordinal(z, y)
        # Las imagenes de referencia deben quedar arriba: si el ajuste ordinal da a la
        # probabilidad de referencia menos peso que el minimo, se sube y el resto se reescala.
        componentes = config["puntaje"]["componentes"]
        minimo = config["puntaje"]["peso_minimo_referencia"]
        i = componentes.index("referencia")
        total = self.pesos.sum()
        if total and self.pesos[i] / total < minimo:
            resto = total - self.pesos[i]
            self.pesos = self.pesos * (1 - minimo) * total / resto if resto else self.pesos
            self.pesos[i] = minimo * total

    def componentes(self, id_global: str) -> list[float]:
        m = self.medidas[id_global]
        valores = [capa_calidad(m, self.config)]
        for nombre in self.config["puntaje"]["componentes"][1:]:
            modelo = self.atributos[nombre]
            valores.append(float(predecir(vector(m, modelo["medidas"], self.estandar)[None, :], np.array(modelo["pesos"]))[0]))
        return valores

    def puntaje(self, id_global: str) -> float:
        z = np.array(self.componentes(id_global))
        return float(100 * z @ self.pesos / self.pesos.sum()) if self.pesos.sum() else 0.0


def estandarizacion(medidas: list[dict[str, str]], nombres: set[str]) -> dict[str, list[float]]:
    """Media y desviacion de cada medida transformada en las imagenes de la fuente de calibracion."""
    salida = {}
    for nombre in sorted(nombres):
        valores = np.array([transformar(nombre, valor_medida(m, nombre)) for m in medidas])
        salida[nombre] = [round(float(valores.mean()), 6), round(float(valores.std() or 1.0), 6)]
    return salida


# ------------------------------------------------------------------ calibracion

def registros_calibracion(raiz: Path, medidas: dict[str, dict[str, str]], config: dict[str, Any]) -> list[dict[str, Any]]:
    """Etiquetas de las fuentes de [calibracion]: todas las humanas o solo las del criterio v2."""
    criterio = config["calibracion"]
    registros = []
    for fuente in criterio["fuentes"]:
        origenes = None if fuente in criterio["fuentes_completas"] else set(criterio["origenes_criterio_v2"])
        registros += etiquetas(raiz, fuente, medidas, config, origenes)
    return registros


def calibrar(raiz: Path) -> tuple[dict[str, Any], list[str]]:
    config = leer_config_aptitud(raiz)
    medidas = indexar(leer_csv(raiz / MEDIDAS), "id_global")
    registros = registros_calibracion(raiz, medidas, config)
    fuentes = config["calibracion"]["fuentes"]
    nombres = {m for atributo in config["atributos"] for m in medidas_atributo(config, atributo)}
    # La estandarizacion usa las imagenes de las tres fuentes: el modelo se aplica a todas.
    estandar = estandarizacion([m for m in medidas.values() if m["fuente"] in fuentes and m["estado_imagen"] == "ok"], nombres)

    # Validacion cruzada estratificada por fuente y clase: cada registro se puntua con
    # modelos que no lo vieron, y cada pliegue contiene las tres colecciones.
    asignacion = pliegues([r["fuente"] + r["clase"] for r in registros])
    fuera = {}
    for k in range(PLIEGUES):
        entrenamiento = [r for r, p in zip(registros, asignacion) if p != k]
        ajuste = Ajuste(config, medidas, entrenamiento, estandar)
        for r, p in zip(registros, asignacion):
            if p == k:
                fuera[r["id_global"]] = {"puntaje": ajuste.puntaje(r["id_global"]),
                                         **dict(zip(config["puntaje"]["componentes"], ajuste.componentes(r["id_global"])))}
    final = Ajuste(config, medidas, registros, estandar)

    umbrales = elegir_umbrales(registros, fuera, config)
    modelo = {
        "fuente_calibracion": "+".join(fuentes),
        "registros": len(registros),
        "registros_por_fuente": {f: sum(r["fuente"] == f for r in registros) for f in fuentes},
        "clases": dict(Counter(r["clase"] for r in registros)),
        "estandarizacion": estandar,
        "atributos": final.atributos,
        "puntaje": {"componentes": config["puntaje"]["componentes"], "pesos": [round(float(w), 6) for w in final.pesos]},
        "umbrales": umbrales,
    }
    informe = construir_informe(registros, fuera, modelo, config)
    return modelo, informe


def elegir_umbrales(registros: list[dict], fuera: dict[str, dict], config: dict[str, Any]) -> dict[str, float]:
    """Umbrales desde los puntajes fuera de pliegue, con los criterios de config [puntaje]."""
    criterio = config["puntaje"]
    puntaje = np.array([fuera[r["id_global"]]["puntaje"] for r in registros])
    util = np.array([r["clase"] in ("principal", "referencia") for r in registros])
    referencia = np.array([r["clase"] == "referencia" for r in registros])
    candidatos = np.unique(np.round(puntaje, 1))

    def primer(condicion):
        for t in candidatos:
            if condicion(t):
                return float(t)
        return float(candidatos[-1])

    principal = primer(lambda t: util[puntaje >= t].mean() >= criterio["precision_principal"])
    descartar = max([float(t) for t in candidatos if util[puntaje < t].sum() <= criterio["principales_bajo_descarte"] * util.sum()] or [0.0])
    # Las imagenes de referencia estan anotadas solo en parte: el umbral se fija por cobertura
    # (recoge esta fraccion de las anotadas) y lo que lo supere sin anotar son candidatas.
    ref = float(np.round(np.quantile(puntaje[referencia], 1 - criterio["cobertura_referencia"]), 1)) if referencia.any() else 100.0
    return {"referencia": max(ref, principal), "principal": principal, "descartar": min(descartar, principal)}


# ------------------------------------------------------------------ informe

def construir_informe(registros: list[dict], fuera: dict[str, dict], modelo: dict, config: dict[str, Any]) -> list[str]:
    clases = [r["clase"] for r in registros]
    puntaje = {r["id_global"]: fuera[r["id_global"]]["puntaje"] for r in registros}

    def p(clase):
        return [puntaje[r["id_global"]] for r in registros if r["clase"] == clase]

    lineas = [
        f"# Calibración de la ficha de aptitud por atributos ({modelo['fuente_calibracion']})",
        "",
        "Generado por `python -m src.imagenes.calibrar_aptitud` (también lo ejecuta `python -m src.flujo`). "
        f"Modelo: `{MODELO}`. Parámetros: `config/aptitud.toml` ([calibracion], [modelo], [atributos] y [puntaje]).",
        "",
        "Calibran todas las decisiones humanas de AIC y, de MET y CMA, solo las tomadas con el criterio v2 "
        "(validación ciega y revisión v2); el resto de MET y CMA se decidió con criterios anteriores.",
        "",
        "Las etiquetas salen del motivo libre del revisor (`src/imagenes/motivos.py`). Se excluyen los descartes "
        "por pertinencia (herramientas, plumaria, muñecos, piezas posteriores a 1532, no andinos) y las fotografías "
        "monocromas: no dependen de lo que la imagen permite identificar. Todas las cifras de esta página son "
        f"**fuera de pliegue** (validación cruzada estratificada de {PLIEGUES} pliegues): cada registro se evalúa "
        "con modelos que no lo vieron.",
        "",
        "## Registros de calibración",
        "",
        "| Clase | " + " | ".join(modelo["registros_por_fuente"]) + " | Total |",
        "|---|" + "---|" * (len(modelo["registros_por_fuente"]) + 1),
        *[f"| {c} | " + " | ".join(str(sum(r["clase"] == c and r["fuente"] == f for r in registros)) for f in modelo["registros_por_fuente"])
          + f" | {clases.count(c)} |" for c in CLASES],
        "",
        "## Atributos estimados desde la imagen",
        "",
        "| Atributo | Etiqueta del revisor (positiva frente a negativa) | Registros | AUC fuera de pliegue | N.º de medidas |",
        "|---|---|---|---|---|",
    ]
    for nombre, definicion in config["atributos"].items():
        usados = [r for r in registros if r[definicion["etiqueta"]] in definicion["positivo"] + definicion["negativo"]]
        pos = [fuera[r["id_global"]][nombre] for r in usados if r[definicion["etiqueta"]] in definicion["positivo"]]
        neg = [fuera[r["id_global"]][nombre] for r in usados if r[definicion["etiqueta"]] in definicion["negativo"]]
        etiqueta = f"{definicion['etiqueta']}: {', '.join(definicion['positivo'])} frente a {', '.join(definicion['negativo'])}"
        lineas.append(f"| {nombre} | {etiqueta} | {len(usados)} | {auc(pos, neg):.3f} | {len(medidas_atributo(config, nombre))} |")

    pesos = modelo["puntaje"]["pesos"]
    total = sum(pesos) or 1
    lineas += [
        "",
        "## Puntaje ordinal",
        "",
        "Media ponderada (pesos no negativos, modelo de odds proporcionales) de la calidad fotográfica y de las "
        "probabilidades de los atributos, escalada a 0-100:",
        "",
        "| Componente | Peso (%) |",
        "|---|---|",
        *[f"| {c} | {100 * w / total:.1f} |" for c, w in zip(modelo["puntaje"]["componentes"], pesos)],
        "",
        "Separación entre clases (AUC del puntaje fuera de pliegue; 0,5 = azar):",
        "",
        "| Comparación | AUC |",
        "|---|---|",
        f"| referencia frente a principal | {auc(p('referencia'), p('principal')):.3f} |",
        f"| principal frente a secundario | {auc(p('principal'), p('secundario')):.3f} |",
        f"| principal frente a descarte por imagen | {auc(p('principal'), p('descarte_imagen')):.3f} |",
        f"| secundario frente a descarte por imagen | {auc(p('secundario'), p('descarte_imagen')):.3f} |",
        f"| referencia o principal frente al resto | {auc(p('referencia') + p('principal'), p('secundario') + p('descarte_imagen')):.3f} |",
        "",
        "Medianas del puntaje por clase: " + ", ".join(f"{c} {np.median(p(c)):.1f}" for c in CLASES if p(c)) + ".",
        "",
        "## Umbrales",
        "",
        f"- Referencia: puntaje ≥ {modelo['umbrales']['referencia']} (recoge el "
        f"{config['puntaje']['cobertura_referencia']:.0%} de las imágenes de referencia del revisor: principal y textil "
        "completo con colores y formas visibles, o marcada como la mejor imagen; las que lo superan sin esa "
        "anotación son candidatas).",
        f"- Principal: puntaje ≥ {modelo['umbrales']['principal']} (al menos el "
        f"{config['puntaje']['precision_principal']:.0%} es principal o referencia).",
        f"- Descartar: puntaje < {modelo['umbrales']['descartar']} (como máximo el "
        f"{config['puntaje']['principales_bajo_descarte']:.0%} de los principales y referencias cae debajo).",
        "- Entre descartar y principal: revisar.",
        "",
    ]
    lineas += [
        "## Separación por fuente",
        "",
        "AUC fuera de pliegue en cada colección (útil = principal o referencia; resto = secundario o descarte "
        "por imagen). Las muestras MET y CMA son pequeñas y reúnen los casos difíciles: sus cifras tienen margen amplio.",
        "",
        "| Fuente | Registros | Útil frente al resto | Colores | Formas | Completitud |",
        "|---|---|---|---|---|---|",
    ]
    for fuente in modelo["registros_por_fuente"]:
        propios = [r for r in registros if r["fuente"] == fuente]
        util = [puntaje[r["id_global"]] for r in propios if r["clase"] in ("principal", "referencia")]
        resto = [puntaje[r["id_global"]] for r in propios if r["clase"] in ("secundario", "descarte_imagen")]
        celdas = [fuente, str(len(propios)), f"{auc(util, resto):.3f}"]
        for nombre in ("colores", "formas", "completitud"):
            definicion = config["atributos"][nombre]
            pos = [fuera[r["id_global"]][nombre] for r in propios if r[definicion["etiqueta"]] in definicion["positivo"]]
            neg = [fuera[r["id_global"]][nombre] for r in propios if r[definicion["etiqueta"]] in definicion["negativo"]]
            celdas.append(f"{auc(pos, neg):.3f}")
        lineas.append("| " + " | ".join(celdas) + " |")
    lineas.append("")
    tabla_bandas = Counter()
    for r in registros:
        s = puntaje[r["id_global"]]
        u = modelo["umbrales"]
        banda = "referencia" if s >= u["referencia"] else "principal" if s >= u["principal"] else "descartar" if s < u["descartar"] else "revisar"
        tabla_bandas[(r["clase"], banda)] += 1
    bandas = ("referencia", "principal", "revisar", "descartar")
    lineas += ["| Clase del revisor \\ banda (fuera de pliegue) | " + " | ".join(bandas) + " |", "|" + "---|" * (len(bandas) + 1)]
    lineas += [f"| {c} | " + " | ".join(str(tabla_bandas[(c, b)]) for b in bandas) + " |" for c in CLASES]
    lineas.append("")
    return lineas


INFORME = "outputs/reports/aptitud_calibracion_v2.md"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Calibra la ficha de aptitud con las decisiones del revisor en AIC, MET y CMA.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    args = parser.parse_args(argv)
    raiz = Path(args.root).resolve()
    modelo, informe = calibrar(raiz)
    (raiz / MODELO).write_bytes((json.dumps(modelo, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    (raiz / INFORME).parent.mkdir(parents=True, exist_ok=True)
    (raiz / INFORME).write_bytes(("\n".join(informe) + "\n").encode("utf-8"))
    print(f"Modelo: {raiz / MODELO} | informe: {raiz / INFORME}")
    print(f"Registros por fuente: {modelo['registros_por_fuente']} | umbrales: {modelo['umbrales']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
