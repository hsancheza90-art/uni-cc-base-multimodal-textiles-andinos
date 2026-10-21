"""Crea o actualiza el workbook de anotacion manual (consolidado v1 o principal v2).

Uso desde la raiz del repositorio (despues de `python -m src.flujo`):

    python -m src.anotacion.workbook                 # v1: data/metadata/anotacion_met_cma_v1.xlsx
    python -m src.anotacion.workbook --corpus v2     # v2: data/metadata/anotacion_v2.xlsx

v1: una fila por registro del consolidado MET+CMA v1, con columnas de contexto (solo lectura)
y columnas de anotacion con listas desplegables del vocabulario controlado.

v2: la hoja "anotacion" tiene las piezas del nucleo de anotacion
(data/metadata/nucleo_anotacion_v2.csv) ordenadas por lote y orden; el resto del principal
v2 va a la hoja "fuera_del_nucleo", con las mismas columnas. Las columnas de anotacion siguen
el orden de las capas de config/anotacion.toml, y los campos multiples tambien muestran la
lista de codigos (aceptan varios separados por "; ").

En los dos casos, si el workbook ya existe se regenera conservando todo lo anotado por
id_global, aunque la pieza haya cambiado de hoja; las anotaciones de registros que ya no
estan en el corpus pasan a la hoja "fuera_del_corpus".
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.anotacion.corpus_v2 import SALIDA_PRINCIPAL_V2
from src.anotacion.nucleo import SALIDA_NUCLEO
from src.consolidado.construir_consolidado import SALIDA_CONSOLIDADO
from src.consolidado.esquema import ATRIBUTOS_ANOTACION, COLUMNAS_ANOTACIONES, CORRECCIONES
from src.utils.archivos import indexar, leer_csv
from src.utils.config import CONFIG_ANOTACION, leer_config_anotacion
from src.utils.libros import leer_hoja
from src.utils.vocabulario import Campo, leer_vocabulario

WORKBOOK = Path("data/metadata/anotacion_met_cma_v1.xlsx")
WORKBOOK_V2 = Path("data/metadata/anotacion_v2.xlsx")
HOJA_ANOTACION = "anotacion"
HOJA_VOCABULARIO = "vocabulario"
HOJA_GUIA = "guia"
HOJA_HUERFANAS = "fuera_del_corpus"
HOJA_FUERA_NUCLEO = "fuera_del_nucleo"

# Columnas con enlace: siempre al final del contexto.
ENLACES = ["url_objeto", "url_imagen", "imagen_local"]
CONTEXTO = [
    "id_global",
    "fuente",
    "decision_curacion_final",
    "grupo_imagen_duplicada",
    "titulo_original",
    "titulo_es_sugerido",
    "cultura",
    "fecha_objeto",
    "tipo_objeto",
    "tipo_superficie",
    "material_normalizado",
    "tecnica_normalizada",
    "estado_imagen",
    *ENLACES,
]
COLUMNAS_EDITABLES = COLUMNAS_ANOTACIONES[1:]
CONTROL_ANOTACION = ["estado_anotacion", "anotador", "fecha_anotacion", "observaciones_anotacion"]

# v2: el contexto de v1 con la posicion en el nucleo y la ficha de aptitud.
CONTEXTO_V2 = [
    "id_global",
    "lote",
    "orden",
    *CONTEXTO[1:-len(ENLACES)],
    "banda_sugerida",
    "referencia_sugerida",
    *ENLACES,
]

# Columna de anotacion -> campo del vocabulario que define sus valores.
CAMPO_DE_COLUMNA: dict[str, str] = {
    **{campo: campo for campo in ATRIBUTOS_ANOTACION},
    **CORRECCIONES,
    "estado_anotacion": "estado_anotacion",
}

RELLENO_CONTEXTO = PatternFill("solid", fgColor="E7E6E6")
RELLENO_ANOTACION = PatternFill("solid", fgColor="FFF2CC")
# v2: un tono por capa en el encabezado, para ubicar la capa de cada columna (solo presentacion).
TONOS_CAPAS = ["FCE4D6", "DDEBF7", "E2EFDA", "EDE2F6", "FFF2CC"]

GUIA = [
    "Anotacion manual del corpus consolidado MET+CMA",
    "",
    "1. Las columnas grises son contexto: no se editan; se regeneran desde el consolidado.",
    "2. Las columnas amarillas se anotan con los codigos de la hoja 'vocabulario' (listas desplegables).",
    "3. Campos multiples (colores, motivos, material_corregido, tecnica_corregida): varios codigos separados por '; '.",
    "4. Las columnas *_corregido solo se llenan para corregir el valor actual; vacias, se conserva el actual.",
    "5. estado_anotacion: si queda vacio y hay anotaciones, se registra 'anotado_manual' (o 'corregido' si hay correcciones).",
    "6. fecha_anotacion con formato AAAA-MM-DD.",
    "7. Despues de anotar: python -m src.flujo valida el workbook y actualiza el consolidado.",
    "",
    "Guia completa: docs/corpus/guia_anotacion_met_cma_v1.md",
]


@dataclass(frozen=True)
class Diseno:
    """Lo que distingue el workbook v1 del v2."""

    workbook: Path
    contexto: list[str]
    editables: list[str]
    guia: list[str]
    hoja_resto: str = ""  # hoja con los registros fuera de la muestra (v2)
    listas_en_multiples: bool = False  # v2: lista de codigos tambien en los campos multiples
    tonos: dict[str, str] | None = None  # v2: columna de anotacion -> tono de su capa

    @property
    def hojas_con_anotaciones(self) -> tuple[str, ...]:
        return tuple(h for h in (HOJA_ANOTACION, self.hoja_resto, HOJA_HUERFANAS) if h)


V1 = Diseno(workbook=WORKBOOK, contexto=CONTEXTO, editables=COLUMNAS_EDITABLES, guia=GUIA)


def leer_capas(config: dict[str, Any]) -> list[dict[str, Any]]:
    """Capas de config/anotacion.toml; deben cubrir exactamente los 12 atributos y las 5 correcciones."""
    capas = config["capas"]
    campos = [campo for capa in capas for campo in capa["campos"]]
    esperados = ATRIBUTOS_ANOTACION + list(CORRECCIONES)
    if sorted(campos) != sorted(esperados) or len(campos) != len(set(campos)):
        raise ValueError(
            f"Las capas de {CONFIG_ANOTACION.as_posix()} deben cubrir una vez cada columna de anotacion: {esperados}"
        )
    return capas


def guia_v2(capas: list[dict[str, Any]]) -> list[str]:
    lineas = [
        "Anotacion estructural del corpus principal v2 (MET, CMA y AIC)",
        "",
        "1. La hoja 'anotacion' tiene las piezas del nucleo, ordenadas por lote y orden: un lote por sesion, en ese orden.",
        "2. La hoja 'fuera_del_nucleo' tiene el resto del principal v2, con las mismas columnas, para anotar mas adelante.",
        "3. Las columnas grises son contexto y no se editan. Las de color se anotan con los codigos de la hoja 'vocabulario'.",
        "4. Se anota sobre la copia de evaluacion (columna imagen_local); url_imagen abre la imagen del museo.",
        "",
        "Capas, en el orden en que se anota cada pieza:",
    ]
    for numero, capa in enumerate(capas, start=1):
        lineas.append(f"   {numero}. {capa['titulo']}: {', '.join(capa['campos'])}")
    lineas += [
        "",
        "Regla de los treinta segundos: si un atributo toma mas de treinta segundos de duda, se deja vacio y la pieza "
        "se marca 'revisar' en estado_anotacion; la duda se explica en observaciones_anotacion.",
        "",
        "Campos multiples (colores, motivos, material_corregido, tecnica_corregida): la lista muestra los codigos; "
        "para registrar varios, escribirlos separados por '; ' (por ejemplo: rojo; crema).",
        "Las columnas *_corregido solo se llenan si el valor gris es incorrecto; vacias, se acepta el valor vigente.",
        "fecha_anotacion con formato AAAA-MM-DD; anotador con nombre o iniciales.",
        "",
        "Protocolo completo: docs/protocolo_anotacion_v2.md",
    ]
    return lineas


def diseno_v2(raiz: Path) -> Diseno:
    capas = leer_capas(leer_config_anotacion(raiz))
    columnas = [campo for capa in capas for campo in capa["campos"]]
    tonos = {campo: TONOS_CAPAS[i % len(TONOS_CAPAS)] for i, capa in enumerate(capas) for campo in capa["campos"]}
    return Diseno(
        workbook=WORKBOOK_V2,
        contexto=CONTEXTO_V2,
        editables=columnas + CONTROL_ANOTACION,
        guia=guia_v2(capas),
        hoja_resto=HOJA_FUERA_NUCLEO,
        listas_en_multiples=True,
        tonos=tonos,
    )


# --------------------------------------------------------------------------- lectura


def leer_anotaciones_existentes(ruta: Path, diseno: Diseno = V1) -> dict[str, dict[str, str]]:
    """Anotaciones del workbook actual (todas sus hojas con registros) por id_global."""
    if not ruta.exists():
        return {}
    anotaciones: dict[str, dict[str, str]] = {}
    for nombre in diseno.hojas_con_anotaciones:
        try:
            filas = leer_hoja(ruta, nombre, clave="id_global")
        except ValueError:
            continue  # las hojas de resto y de huerfanas solo existen si tienen registros
        for datos in filas:
            if any(datos.get(c) for c in diseno.editables):
                anotaciones[datos["id_global"]] = {c: datos.get(c, "") for c in diseno.editables}
    return anotaciones


def registros_v2(raiz: Path) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Piezas del nucleo (por lote y orden) y resto del principal v2 (en el orden del corpus)."""
    corpus = leer_csv(raiz / SALIDA_PRINCIPAL_V2, recortar=False)
    nucleo = indexar(leer_csv(raiz / SALIDA_NUCLEO), "id_global")
    faltan = sorted(set(nucleo) - {fila["id_global"] for fila in corpus})
    if faltan:
        raise ValueError(f"Piezas del nucleo que no estan en el principal v2: {faltan[:10]}")
    dentro, fuera = [], []
    for fila in corpus:
        posicion = nucleo.get(fila["id_global"])
        if posicion:
            dentro.append({**fila, "lote": int(posicion["lote"]), "orden": int(posicion["orden"])})
        else:
            fuera.append({**fila, "lote": "", "orden": ""})
    dentro.sort(key=lambda fila: (fila["lote"], fila["orden"]))
    return dentro, fuera


# --------------------------------------------------------------------------- escritura


def escribir_hoja_vocabulario(libro: Workbook, vocabulario: dict[str, Campo]) -> dict[str, str]:
    """Una pareja de columnas (codigo, etiqueta) por campo; devuelve el rango de codigos de cada campo."""
    hoja = libro.create_sheet(HOJA_VOCABULARIO)
    rangos: dict[str, str] = {}
    columna = 1
    for campo in vocabulario.values():
        if not campo.anotable:
            continue
        letra_codigo, letra_etiqueta = get_column_letter(columna), get_column_letter(columna + 1)
        hoja[f"{letra_codigo}1"] = campo.nombre
        hoja[f"{letra_etiqueta}1"] = "etiqueta"
        hoja[f"{letra_codigo}1"].font = hoja[f"{letra_etiqueta}1"].font = Font(bold=True)
        for fila, (codigo, etiqueta) in enumerate(campo.etiquetas.items(), start=2):
            hoja[f"{letra_codigo}{fila}"] = codigo
            hoja[f"{letra_etiqueta}{fila}"] = etiqueta
        rangos[campo.nombre] = f"{HOJA_VOCABULARIO}!${letra_codigo}$2:${letra_codigo}${len(campo.etiquetas) + 1}"
        hoja.column_dimensions[letra_codigo].width = 28
        hoja.column_dimensions[letra_etiqueta].width = 28
        columna += 3
    return rangos


def agregar_validaciones(
    hoja, vocabulario: dict[str, Campo], rangos: dict[str, str], filas: int, listas_en_multiples: bool = False
) -> None:
    if not filas:
        return
    encabezados = [celda.value for celda in hoja[1]]
    for columna, nombre_campo in CAMPO_DE_COLUMNA.items():
        campo = vocabulario[nombre_campo]
        letra = get_column_letter(encabezados.index(columna) + 1)
        rango_celdas = f"{letra}2:{letra}{filas + 1}"
        if campo.multiple and listas_en_multiples:
            # La lista muestra los codigos; sin mensaje de error, para admitir varios separados por "; ".
            validacion = DataValidation(type="list", formula1=f"={rangos[nombre_campo]}", allow_blank=True)
            validacion.showErrorMessage = False
            validacion.promptTitle = columna[:32]
            validacion.prompt = "Elija un codigo o escriba varios separados por '; '."
            validacion.showInputMessage = True
        elif campo.multiple:
            validacion = DataValidation(allow_blank=True)
            validacion.promptTitle = columna[:32]
            validacion.prompt = ("Codigos separados por '; ': " + ", ".join(campo.codigos))[:255]
            validacion.showInputMessage = True
        else:
            validacion = DataValidation(type="list", formula1=f"={rangos[nombre_campo]}", allow_blank=True)
            validacion.error = f"Use un codigo de la hoja '{HOJA_VOCABULARIO}'."
            validacion.showErrorMessage = True
        validacion.add(rango_celdas)
        hoja.add_data_validation(validacion)


def escribir_enlace(celda, url: str) -> None:
    if url:
        celda.value = url
        celda.hyperlink = url
        celda.style = "Hyperlink"


def escribir_registros(hoja, registros: list[dict[str, Any]], anotaciones: dict[str, dict[str, str]], diseno: Diseno) -> None:
    """Encabezado, una fila por registro (contexto y anotaciones previas), anchos, paneles y filtro."""
    encabezados = diseno.contexto + diseno.editables
    hoja.append(encabezados)
    for indice, nombre in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=indice)
        celda.font = Font(bold=True)
        if nombre in diseno.contexto:
            celda.fill = RELLENO_CONTEXTO
        elif diseno.tonos and nombre in diseno.tonos:
            celda.fill = PatternFill("solid", fgColor=diseno.tonos[nombre])
        else:
            celda.fill = RELLENO_ANOTACION
        celda.alignment = Alignment(wrap_text=True, vertical="top")

    columna_de = {nombre: indice for indice, nombre in enumerate(encabezados, start=1)}
    for numero, registro in enumerate(registros, start=2):
        for nombre in diseno.contexto[:-len(ENLACES)]:
            hoja.cell(row=numero, column=columna_de[nombre], value=registro.get(nombre, ""))
        escribir_enlace(hoja.cell(row=numero, column=columna_de["url_objeto"]), registro["url_objeto"])
        escribir_enlace(hoja.cell(row=numero, column=columna_de["url_imagen"]), registro["url_imagen"])
        if registro.get("ruta_imagen_local"):
            # El workbook vive en data/metadata; la imagen local, en data/images.
            relativa = "../" + registro["ruta_imagen_local"].removeprefix("data/")
            escribir_enlace(hoja.cell(row=numero, column=columna_de["imagen_local"]), relativa)
        previo = anotaciones.get(registro["id_global"], {})
        for nombre in diseno.editables:
            if previo.get(nombre):
                hoja.cell(row=numero, column=columna_de[nombre], value=previo[nombre])

    anchos = {"titulo_original": 34, "titulo_es_sugerido": 22, "cultura": 24, "url_objeto": 22, "url_imagen": 22,
              "lote": 6, "orden": 7}
    for nombre, indice in columna_de.items():
        hoja.column_dimensions[get_column_letter(indice)].width = anchos.get(nombre, 18)
    hoja.freeze_panes = "C2" if diseno.contexto[1] != "lote" else "D2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{len(registros) + 1}"


def construir_libro(
    registros: list[dict[str, Any]],
    anotaciones: dict[str, dict[str, str]],
    vocabulario: dict[str, Campo],
    diseno: Diseno = V1,
    resto: list[dict[str, Any]] | None = None,
) -> Workbook:
    resto = resto or []
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_ANOTACION
    escribir_registros(hoja, registros, anotaciones, diseno)
    hojas = [(hoja, len(registros))]
    if diseno.hoja_resto:
        hoja_resto = libro.create_sheet(diseno.hoja_resto)
        escribir_registros(hoja_resto, resto, anotaciones, diseno)
        hojas.append((hoja_resto, len(resto)))

    rangos = escribir_hoja_vocabulario(libro, vocabulario)
    for hoja_datos, filas in hojas:
        agregar_validaciones(hoja_datos, vocabulario, rangos, filas, diseno.listas_en_multiples)

    guia = libro.create_sheet(HOJA_GUIA, 0)
    for numero, linea in enumerate(diseno.guia, start=1):
        guia.cell(row=numero, column=1, value=linea)
    guia["A1"].font = Font(bold=True, size=14)
    guia.column_dimensions["A"].width = 120

    ids_actuales = {registro["id_global"] for registro in registros + resto}
    huerfanas = {id_: datos for id_, datos in anotaciones.items() if id_ not in ids_actuales}
    if huerfanas:
        hoja_h = libro.create_sheet(HOJA_HUERFANAS)
        hoja_h.append(["id_global"] + diseno.editables)
        for id_global, datos in sorted(huerfanas.items()):
            hoja_h.append([id_global] + [datos.get(c, "") for c in diseno.editables])
    return libro


# --------------------------------------------------------------------------- principal


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Crea o actualiza el workbook de anotacion manual.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    parser.add_argument(
        "--corpus", choices=["v1", "v2"], default="v1",
        help="v1: consolidado MET+CMA v1 (por defecto); v2: nucleo y resto del principal v2.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    raiz = Path(args.root).resolve()
    if args.corpus == "v2":
        diseno = diseno_v2(raiz)
        registros, resto = registros_v2(raiz)
    else:
        diseno = V1
        registros, resto = leer_csv(raiz / SALIDA_CONSOLIDADO, recortar=False), []
    ruta = raiz / diseno.workbook
    anotaciones = leer_anotaciones_existentes(ruta, diseno)
    libro = construir_libro(registros, anotaciones, leer_vocabulario(raiz), diseno, resto)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        libro.save(ruta)
    except PermissionError:
        print(f"No se pudo escribir {ruta}: cierre el archivo en Excel y vuelva a intentarlo.")
        return 1
    ids = {r["id_global"] for r in registros + resto}
    conservadas = sum(1 for id_ in anotaciones if id_ in ids)
    print(f"Workbook de anotacion: {ruta}")
    detalle = f" (+ {len(resto)} en {diseno.hoja_resto})" if diseno.hoja_resto else ""
    print(f"Registros: {len(registros)}{detalle} | anotaciones conservadas: {conservadas} | fuera del corpus: {len(anotaciones) - conservadas}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
