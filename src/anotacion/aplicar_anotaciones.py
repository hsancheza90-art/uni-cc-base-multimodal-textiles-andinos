"""Valida el workbook de anotacion y exporta las anotaciones a CSV.

Paso del flujo (sin red). Lee data/metadata/anotacion_met_cma_v1.xlsx, recodifica
etiquetas a codigos, valida cada valor contra config/vocabulario.toml y escribe:

- data/metadata/anotaciones_met_cma_v1.csv: solo registros con alguna anotacion.
- outputs/reports/anotacion_met_cma_v1_validacion.md: avance y errores.

Si hay valores fuera del vocabulario o fechas invalidas, termina con codigo 1 y no
actualiza el CSV, para que el consolidado nunca incorpore anotaciones invalidas.
Sin workbook, no hace nada.
"""

from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from src.anotacion.workbook import CAMPO_DE_COLUMNA, COLUMNAS_EDITABLES, HOJA_ANOTACION, WORKBOOK, texto_celda
from src.consolidado.construir_consolidado import (
    ANOTACIONES,
    clave_orden,
    construir_registros_base,
    escribir_csv,
    escribir_texto,
)
from src.consolidado.esquema import ATRIBUTOS_ANOTACION, COLUMNAS_ANOTACIONES, CORRECCIONES
from src.utils.vocabulario import Campo, leer_vocabulario

REPORTE = "outputs/reports/anotacion_met_cma_v1_validacion.md"
FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def leer_workbook(ruta: Path) -> list[dict[str, str]]:
    hoja = load_workbook(ruta, data_only=True)[HOJA_ANOTACION]
    encabezados = [texto_celda(celda.value) for celda in hoja[1]]
    filas = []
    for valores in hoja.iter_rows(min_row=2, values_only=True):
        fila = {encabezados[i]: texto_celda(v) for i, v in enumerate(valores) if i < len(encabezados)}
        if fila.get("id_global"):
            filas.append(fila)
    return filas


def normalizar_fila(fila: dict[str, str], vocabulario: dict[str, Campo]) -> tuple[dict[str, str] | None, list[str]]:
    """Devuelve la anotacion recodificada (o None si no hay nada anotado) y sus errores."""
    if not any(fila.get(columna) for columna in COLUMNAS_EDITABLES):
        return None, []
    errores: list[str] = []
    salida = {"id_global": fila["id_global"]}
    for columna in COLUMNAS_EDITABLES:
        valor = fila.get(columna, "")
        if columna in CAMPO_DE_COLUMNA and valor:
            campo = vocabulario[CAMPO_DE_COLUMNA[columna]]
            valor = campo.recodificar(valor)
            invalidos = campo.invalidos(valor)
            if invalidos:
                errores.append(f"{fila['id_global']}: {columna} fuera del vocabulario: {invalidos}")
        salida[columna] = valor

    if salida["fecha_anotacion"] and not FECHA.match(salida["fecha_anotacion"]):
        errores.append(f"{fila['id_global']}: fecha_anotacion debe tener formato AAAA-MM-DD")
    if not salida["estado_anotacion"] or salida["estado_anotacion"] == "sin_anotar":
        hay_correccion = any(salida[c] for c in CORRECCIONES)
        hay_atributos = any(salida[c] for c in ATRIBUTOS_ANOTACION)
        if hay_correccion:
            salida["estado_anotacion"] = "corregido"
        elif hay_atributos:
            salida["estado_anotacion"] = "anotado_manual"
    return salida, errores


def lineas_reporte(total: int, anotaciones: list[dict[str, str]], errores: list[str], huerfanas: list[str]) -> list[str]:
    estados = Counter(fila["estado_anotacion"] or "sin_anotar" for fila in anotaciones)
    estados["sin_anotar"] += total - len(anotaciones)
    lineas = [
        "# Validacion de la anotacion manual MET+CMA v1",
        "",
        f"Workbook: `{WORKBOOK.as_posix()}`. Generado por `src/anotacion/aplicar_anotaciones.py`.",
        "",
        f"- Registros del consolidado: {total}",
        f"- Registros con anotacion: {len(anotaciones)}",
        f"- Errores de validacion: {len(errores)}",
        "",
        "## Estado de anotacion",
        "",
        "| estado_anotacion | Registros |",
        "|---|---:|",
    ]
    lineas += [f"| {estado} | {n} |" for estado, n in sorted(estados.items())]
    lineas += ["", "## Cobertura por atributo", "", "| Atributo | Registros con valor |", "|---|---:|"]
    lineas += [f"| `{c}` | {sum(bool(f[c]) for f in anotaciones)} |" for c in ATRIBUTOS_ANOTACION + list(CORRECCIONES)]
    if huerfanas:
        lineas += ["", "## Anotaciones de registros fuera del consolidado", ""]
        lineas += [f"- {id_}" for id_ in huerfanas]
    if errores:
        lineas += ["", "## Errores", ""]
        lineas += [f"- {error}" for error in errores]
    return lineas


def aplicar(raiz: Path) -> int:
    ruta = raiz / WORKBOOK
    if not ruta.exists():
        print(f"No existe {WORKBOOK.as_posix()}; no hay anotaciones que aplicar.")
        return 0

    vocabulario = leer_vocabulario(raiz)
    registros = construir_registros_base(raiz)
    orden = {fila["id_global"]: clave_orden(fila) for fila in registros}

    anotaciones: list[dict[str, str]] = []
    errores: list[str] = []
    for fila in leer_workbook(ruta):
        anotacion, errores_fila = normalizar_fila(fila, vocabulario)
        errores += errores_fila
        if anotacion:
            anotaciones.append(anotacion)
    huerfanas = sorted(a["id_global"] for a in anotaciones if a["id_global"] not in orden)
    anotaciones.sort(key=lambda a: orden.get(a["id_global"], (9, 0, a["id_global"])))

    escribir_texto(raiz / REPORTE, lineas_reporte(len(registros), anotaciones, errores, huerfanas))
    if errores:
        print(f"Anotaciones con {len(errores)} errores; no se actualizo {ANOTACIONES}. Ver {REPORTE}.")
        for error in errores[:20]:
            print(f"- {error}")
        return 1

    escribir_csv(raiz / ANOTACIONES, COLUMNAS_ANOTACIONES, anotaciones)
    print(f"Anotaciones validas: {len(anotaciones)} de {len(registros)} registros -> {ANOTACIONES}")
    return 0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Valida el workbook de anotacion y exporta las anotaciones.")
    parser.add_argument("--root", default=".", help="Raiz del repositorio.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    return aplicar(Path(parse_args(argv).root).resolve())


if __name__ == "__main__":
    raise SystemExit(main())
