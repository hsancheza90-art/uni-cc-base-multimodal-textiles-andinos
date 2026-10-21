"""Pruebas del corpus principal v2, el nucleo de anotacion, el workbook v2 y el acuerdo entre anotadores."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.anotacion import acuerdo, corpus_v2, nucleo, workbook
from src.imagenes.aptitud import SALIDA_APTITUD
from src.imagenes.descargar_evaluacion import MANIFIESTO_EVALUACION
from src.utils.archivos import escribir_csv, leer_csv_con_columnas
from src.utils.config import RAIZ_REPO, leer_config, leer_config_anotacion
from src.utils.fuentes import FUENTES
from src.utils.metricas import jaccard, kappa_detalle
from src.utils.vocabulario import leer_vocabulario
from tests.insumos import copiar_insumos
from tests.libros import contenido_libro


def insumos_corpus_v2() -> list[str]:
    """Archivos versionados que lee src/anotacion/corpus_v2.py."""
    base_met = [d["ruta"] for d in leer_config()["met"]["base"].values() if isinstance(d, dict)]
    return [
        "config/corpus.toml",
        "config/vocabulario.toml",
        *(fuente.salidas["principal"] for fuente in FUENTES.values()),
        *(fuente.candidatos for fuente in FUENTES.values() if fuente.candidatos),
        *base_met,
        MANIFIESTO_EVALUACION,
        SALIDA_APTITUD,
    ]


def leer(ruta: Path) -> list[dict[str, str]]:
    with ruta.open(newline="", encoding="utf-8") as archivo:
        return list(csv.DictReader(archivo))


@pytest.fixture()
def repo_v2(tmp_path: Path) -> Path:
    copiar_insumos(tmp_path, insumos_corpus_v2())
    return tmp_path


# --------------------------------------------------------------------------- corpus principal v2


def test_corpus_v2_se_regenera_identico(repo_v2: Path) -> None:
    assert corpus_v2.main(["--root", str(repo_v2)]) == 0
    for relativa in (corpus_v2.SALIDA_PRINCIPAL_V2, corpus_v2.SALIDA_REPORTE):
        assert (repo_v2 / relativa).read_bytes() == (RAIZ_REPO / relativa).read_bytes(), relativa


def test_corpus_v2_une_las_tres_fuentes_sin_colisiones() -> None:
    filas = leer(RAIZ_REPO / corpus_v2.SALIDA_PRINCIPAL_V2)
    esperados = sum(len(leer(RAIZ_REPO / f.salidas["principal"])) for f in FUENTES.values())
    ids = [fila["id_global"] for fila in filas]
    assert len(filas) == esperados
    assert len(set(ids)) == len(ids)
    assert all(fila["id_global"] == f"{fila['fuente']}:{fila['id_fuente']}" for fila in filas)
    assert {fila["decision_curacion_final"] for fila in filas} == {"principal"}


def test_corpus_v2_usa_el_esquema_armonizado_y_el_vocabulario() -> None:
    columnas, filas = leer_csv_con_columnas(RAIZ_REPO / corpus_v2.SALIDA_PRINCIPAL_V2)
    assert columnas == corpus_v2.COLUMNAS_V2
    vocabulario = leer_vocabulario()
    for campo in vocabulario.values():
        if campo.nombre in columnas:
            assert not [f["id_global"] for f in filas if campo.invalidos(f[campo.nombre])], campo.nombre
    assert {f["estado_anotacion"] for f in filas} == {"sin_anotar"}


def test_corpus_v2_falla_si_un_id_global_se_repite(repo_v2: Path) -> None:
    ruta = repo_v2 / FUENTES["CMA"].salidas["principal"]
    columnas, filas = leer_csv_con_columnas(ruta)
    escribir_csv(ruta, columnas, filas + [filas[0]])
    assert corpus_v2.main(["--root", str(repo_v2)]) == 1
    assert not (repo_v2 / corpus_v2.SALIDA_PRINCIPAL_V2).exists()


def test_corpus_v2_conserva_registros_sin_ficha_de_aptitud(repo_v2: Path) -> None:
    ruta = repo_v2 / SALIDA_APTITUD
    columnas, filas = leer_csv_con_columnas(ruta)
    quitado = "AIC:" + leer(repo_v2 / FUENTES["AIC"].salidas["principal"])[0]["id_fuente"]
    escribir_csv(ruta, columnas, [f for f in filas if f["id_global"] != quitado])

    assert corpus_v2.main(["--root", str(repo_v2)]) == 0
    salida = {f["id_global"]: f for f in leer(repo_v2 / corpus_v2.SALIDA_PRINCIPAL_V2)}
    assert quitado in salida
    assert all(salida[quitado][c] == "" for c in corpus_v2.COLUMNAS_APTITUD)
    reporte = (repo_v2 / corpus_v2.SALIDA_REPORTE).read_text(encoding="utf-8")
    assert f"sin fila en `{SALIDA_APTITUD}`: 1 ({quitado})" in reporte


def test_valor_curado_fuera_del_vocabulario_lo_reemplaza_la_regla() -> None:
    vocabulario = leer_vocabulario()
    fila = {campo: "" for campo in corpus_v2.COLUMNAS}
    fila.update({"fuente": "CMA", "tipo_objeto": "Textile", "material_normalizado": "algodón"})
    descartados = corpus_v2.descartar_fuera_de_vocabulario(fila, vocabulario)
    assert descartados == [("CMA", "tipo_objeto", "Textile")]
    assert fila["tipo_objeto"] == ""
    assert fila["material_normalizado"] == "algodón"  # etiqueta valida: se recodifica despues


# --------------------------------------------------------------------------- nucleo de anotacion


@pytest.fixture()
def repo_nucleo(tmp_path: Path) -> Path:
    copiar_insumos(tmp_path, ["config/anotacion.toml", corpus_v2.SALIDA_PRINCIPAL_V2])
    return tmp_path


def correr_nucleo(raiz: Path, semilla: int | None = None) -> nucleo.Nucleo:
    config = leer_config_anotacion(raiz)
    registros = leer(raiz / corpus_v2.SALIDA_PRINCIPAL_V2)
    return nucleo.construir_nucleo(registros, config, int(config["nucleo"]["semilla"]) if semilla is None else semilla)


def test_nucleo_se_regenera_byte_a_byte(repo_nucleo: Path) -> None:
    for _ in range(2):
        assert nucleo.main(["--root", str(repo_nucleo)]) == 0
        for relativa in (nucleo.SALIDA_NUCLEO, nucleo.SALIDA_REPORTE):
            assert (repo_nucleo / relativa).read_bytes() == (RAIZ_REPO / relativa).read_bytes(), relativa


def test_cuotas_por_fuente_suman_el_tamano_y_respetan_la_proporcion() -> None:
    resultado = correr_nucleo(RAIZ_REPO)
    tamano = int(leer_config_anotacion()["nucleo"]["tamano"])
    total = sum(resultado.disponibles_fuente.values())
    assert sum(resultado.cuotas_fuente.values()) == len(resultado.filas) == tamano
    for fuente, cuota in resultado.cuotas_fuente.items():
        assert abs(cuota - tamano * resultado.disponibles_fuente[fuente] / total) < 1
        assert sum(f["fuente"] == fuente for f in resultado.filas) == cuota


def test_otra_semilla_cambia_las_piezas_pero_no_las_cuotas() -> None:
    a, b = correr_nucleo(RAIZ_REPO), correr_nucleo(RAIZ_REPO, semilla=7)
    assert {f["id_global"] for f in a.filas} != {f["id_global"] for f in b.filas}
    cuotas = lambda r: [(e.fuente, e.cultura, e.banda, e.cuota, e.asignadas) for e in r.estratos]  # noqa: E731
    assert cuotas(a) == cuotas(b)


def test_mayores_restos_da_la_unidad_sobrante_al_mayor_resto() -> None:
    # Cuotas exactas 1,6 / 1,2 / 0,2: las partes enteras suman 2 y la unidad restante va al mayor resto (a).
    assert nucleo.mayores_restos(3, {"a": 8, "b": 6, "c": 1}) == {"a": 2, "b": 1, "c": 0}
    # Un estrato pequeño entra si su resto lo justifica: exactas 1,54 / 1,54 / 0,92; c tiene el mayor
    # resto y entra con 1; la segunda unidad desempata a favor de a (mismo peso, primera clave).
    assert nucleo.mayores_restos(4, {"a": 15, "b": 15, "c": 9}) == {"a": 2, "b": 1, "c": 1}


def test_estrato_con_menos_piezas_que_su_cuota_redistribuye_el_faltante() -> None:
    # a pide 5 y tiene 3: los 2 que faltan van a b (3 libres) y c (1 libre) en proporcion a lo libre,
    # exactas 1,5 / 0,5; el empate de restos se resuelve por mayor peso (b).
    asignadas, faltantes = nucleo.ajustar_a_disponibles({"a": 5, "b": 2, "c": 3}, {"a": 3, "b": 5, "c": 4})
    assert asignadas == {"a": 3, "b": 4, "c": 3}
    assert faltantes == {"a": 2, "b": 0, "c": 0}


def test_el_informe_registra_los_faltantes() -> None:
    resultado = correr_nucleo(RAIZ_REPO)
    resultado.estratos[0].faltantes = 2
    lineas = nucleo.lineas_reporte(resultado, leer_config_anotacion())
    assert "Faltantes redistribuidos dentro de su fuente: 2 pieza(s)." in lineas


def test_ningun_lote_supera_el_tope_de_una_fuente() -> None:
    config = leer_config_anotacion()
    tope, tamano_lote = float(config["nucleo"]["tope_fuente_lote"]), int(config["nucleo"]["tamano_lote"])
    filas = leer(RAIZ_REPO / nucleo.SALIDA_NUCLEO)
    for lote in {f["lote"] for f in filas}:
        propias = [f for f in filas if f["lote"] == lote]
        assert len(propias) == tamano_lote
        assert sorted(int(f["orden"]) for f in propias) == list(range(1, tamano_lote + 1))
        assert max(Counter(f["fuente"] for f in propias).values()) <= tope * tamano_lote


def test_un_tope_imposible_falla_con_mensaje_claro(repo_nucleo: Path, capsys: pytest.CaptureFixture[str]) -> None:
    ruta = repo_nucleo / "config/anotacion.toml"
    ruta.write_text(ruta.read_text(encoding="utf-8").replace("tope_fuente_lote = 0.6", "tope_fuente_lote = 0.4"), encoding="utf-8")
    assert nucleo.main(["--root", str(repo_nucleo)]) == 1
    assert "tope 40%" in capsys.readouterr().out
    assert not (repo_nucleo / nucleo.SALIDA_NUCLEO).exists()


def test_cultura_agrupada_usa_terminos_casos_y_sin_cultura() -> None:
    config = leer_config_anotacion()
    casos = nucleo.leer_casos(config)
    assert nucleo.agrupar_cultura("Peru, South Coast, Wari Culture", config, casos) == "wari"
    assert nucleo.agrupar_cultura("huari", config, casos) == "wari"
    assert nucleo.agrupar_cultura("Nasca-Wari", config, casos) == "wari"
    assert nucleo.agrupar_cultura("Lambayeque", config, casos) == "otras"
    assert nucleo.agrupar_cultura("", config, casos) == "sin_cultura"
    with pytest.raises(nucleo.ErrorNucleo, match="nombra varios grupos"):
        nucleo.agrupar_cultura("Paracas-Wari", config, casos)


# --------------------------------------------------------------------------- workbook v2


@pytest.fixture()
def repo_workbook(tmp_path: Path) -> Path:
    copiar_insumos(tmp_path, [
        "config/anotacion.toml", "config/vocabulario.toml", corpus_v2.SALIDA_PRINCIPAL_V2, nucleo.SALIDA_NUCLEO,
    ])
    assert workbook.main(["--root", str(tmp_path), "--corpus", "v2"]) == 0
    return tmp_path


def filas_hoja(ruta: Path, hoja: str) -> list[dict[str, Any]]:
    datos = load_workbook(ruta)[hoja]
    encabezados = [celda.value for celda in datos[1]]
    return [dict(zip(encabezados, fila)) for fila in datos.iter_rows(min_row=2, values_only=True)]


def anotar_v2(ruta: Path, valores: dict[str, dict[str, str]]) -> None:
    libro = load_workbook(ruta)
    for hoja in (workbook.HOJA_ANOTACION, workbook.HOJA_FUERA_NUCLEO):
        datos = libro[hoja]
        encabezados = [celda.value for celda in datos[1]]
        for fila in datos.iter_rows(min_row=2):
            for columna, valor in valores.get(fila[0].value, {}).items():
                fila[encabezados.index(columna)].value = valor
    libro.save(ruta)


def test_workbook_v2_tiene_el_nucleo_ordenado_y_17_columnas_vacias_con_lista(repo_workbook: Path) -> None:
    ruta = repo_workbook / workbook.WORKBOOK_V2
    libro = load_workbook(ruta)
    assert libro.sheetnames == ["guia", "anotacion", "fuera_del_nucleo", "vocabulario"]
    hoja = libro[workbook.HOJA_ANOTACION]
    encabezados = [celda.value for celda in hoja[1]]
    anotables = [c for c in encabezados if c in workbook.CAMPO_DE_COLUMNA and c != "estado_anotacion"]
    assert len(anotables) == 17
    assert anotables == [campo for capa in leer_config_anotacion()["capas"] for campo in capa["campos"]]
    columnas_con_lista = {
        str(v.sqref).split(":")[0].rstrip("0123456789") for v in hoja.data_validations.dataValidation if v.type == "list"
    }
    assert all(get_column_letter(encabezados.index(c) + 1) in columnas_con_lista for c in anotables)

    filas = filas_hoja(ruta, workbook.HOJA_ANOTACION)
    nucleo_csv = leer(RAIZ_REPO / nucleo.SALIDA_NUCLEO)
    assert [f["id_global"] for f in filas] == [f["id_global"] for f in nucleo_csv]
    assert [(f["lote"], f["orden"]) for f in filas] == sorted((f["lote"], f["orden"]) for f in filas)
    assert all(f[c] is None for f in filas for c in anotables)
    resto = filas_hoja(ruta, workbook.HOJA_FUERA_NUCLEO)
    assert len(filas) + len(resto) == len(leer(RAIZ_REPO / corpus_v2.SALIDA_PRINCIPAL_V2))

    guia = [fila[0] for fila in libro["guia"].iter_rows(values_only=True) if fila[0]]
    assert any("treinta segundos" in linea for linea in guia)
    capas = [linea.strip() for linea in guia if linea.strip()[:2] in {"1.", "2.", "3.", "4.", "5."} and ":" in linea]
    assert [linea.split(":")[0][3:] for linea in capas[-5:]] == [c["titulo"] for c in leer_config_anotacion()["capas"]]


def test_workbook_v2_versionado_coincide_con_el_generado(repo_workbook: Path) -> None:
    assert contenido_libro(repo_workbook / workbook.WORKBOOK_V2) == contenido_libro(RAIZ_REPO / workbook.WORKBOOK_V2)


def test_workbook_v2_conserva_anotaciones_aunque_la_pieza_cambie_de_hoja(repo_workbook: Path) -> None:
    ruta = repo_workbook / workbook.WORKBOOK_V2
    antes = {f["id_global"] for f in leer(repo_workbook / nucleo.SALIDA_NUCLEO)}
    despues = {f["id_global"] for f in correr_nucleo(repo_workbook, semilla=7).filas}
    entra = sorted(despues - antes)[0]  # estaba fuera del nucleo y entra
    sale = sorted(antes - despues)[0]   # estaba en el nucleo y sale
    anotar_v2(ruta, {
        entra: {"composicion": "bandas", "colores": "rojo; crema", "observaciones_anotacion": "franjas"},
        sale: {"simetria": "bilateral", "estado_anotacion": "revisar"},
    })

    assert nucleo.main(["--root", str(repo_workbook), "--semilla", "7"]) == 0
    assert workbook.main(["--root", str(repo_workbook), "--corpus", "v2"]) == 0

    dentro = {f["id_global"]: f for f in filas_hoja(ruta, workbook.HOJA_ANOTACION)}
    fuera = {f["id_global"]: f for f in filas_hoja(ruta, workbook.HOJA_FUERA_NUCLEO)}
    assert dentro[entra]["composicion"] == "bandas" and dentro[entra]["colores"] == "rojo; crema"
    assert dentro[entra]["observaciones_anotacion"] == "franjas"
    assert fuera[sale]["simetria"] == "bilateral" and fuera[sale]["estado_anotacion"] == "revisar"
    assert workbook.HOJA_HUERFANAS not in load_workbook(ruta).sheetnames


def test_workbook_v1_no_cambia(tmp_path: Path) -> None:
    copiar_insumos(tmp_path, [
        "config/vocabulario.toml", "data/processed/corpus_textiles_andinos_met_cma_v1_consolidado.csv",
    ])
    # Sin workbook previo y regenerando sobre el versionado: los dos coinciden con el actual.
    assert workbook.main(["--root", str(tmp_path)]) == 0
    assert contenido_libro(tmp_path / workbook.WORKBOOK) == contenido_libro(RAIZ_REPO / workbook.WORKBOOK)
    copiar_insumos(tmp_path, [str(workbook.WORKBOOK.as_posix())])
    assert workbook.main(["--root", str(tmp_path), "--corpus", "v1"]) == 0
    assert contenido_libro(tmp_path / workbook.WORKBOOK) == contenido_libro(RAIZ_REPO / workbook.WORKBOOK)


# --------------------------------------------------------------------------- acuerdo entre anotadores


def test_kappa_de_una_tabla_hecha_a_mano() -> None:
    # 50 piezas: 20 si/si, 5 si/no, 10 no/si, 15 no/no. Observado 0,70; esperado 0,5*0,6 + 0,5*0,4 = 0,50.
    pares = [("si", "si")] * 20 + [("si", "no")] * 5 + [("no", "si")] * 10 + [("no", "no")] * 15
    resultado = kappa_detalle(pares)
    assert resultado.n == 50
    assert resultado.observado == pytest.approx(0.70)
    assert resultado.esperado == pytest.approx(0.50)
    assert resultado.kappa == pytest.approx(0.40)


def test_kappa_con_acuerdo_perfecto_es_uno() -> None:
    pares = [("bandas", "bandas"), ("paneles", "paneles"), ("reticula", "reticula"), ("bandas", "bandas")]
    assert kappa_detalle(pares).kappa == pytest.approx(1.0)


def test_kappa_con_acuerdo_al_azar_es_cercano_a_cero() -> None:
    generador = np.random.default_rng(20261020)
    categorias = ["bandas", "paneles", "reticula", "composicion_libre"]
    a, b = generador.choice(categorias, 4000), generador.choice(categorias, 4000)
    assert abs(kappa_detalle(list(zip(a, b))).kappa) < 0.05


def test_kappa_ponderado_supera_al_simple_si_los_desacuerdos_son_vecinos() -> None:
    orden = ["bajo", "medio", "alto"]
    pares = [("bajo", "bajo"), ("medio", "medio"), ("alto", "alto"), ("bajo", "medio"), ("medio", "alto"), ("alto", "medio")]
    simple, ponderado = kappa_detalle(pares, orden), kappa_detalle(pares, orden, ponderado=True)
    assert simple.kappa == pytest.approx(0.25)
    assert ponderado.kappa == pytest.approx(0.40)
    assert ponderado.kappa > simple.kappa


def test_atributo_constante_o_sin_piezas_no_es_estimable() -> None:
    constante = kappa_detalle([("si", "si")] * 8, ["no", "si"])
    assert constante.kappa is None and "una sola categoria" in constante.motivo
    vacio = kappa_detalle([])
    assert vacio.kappa is None and vacio.motivo == "sin piezas comparables"
    assert acuerdo.alcanza(constante, 0.6).startswith("no estimable")


def test_jaccard_por_pieza() -> None:
    assert jaccard({"rojo", "crema"}, {"rojo", "crema"}) == 1.0
    assert jaccard({"rojo", "crema"}, {"rojo", "negro"}) == pytest.approx(1 / 3)


def test_informe_de_acuerdo_versionado_se_regenera(tmp_path: Path) -> None:
    copiar_insumos(tmp_path, ["config/anotacion.toml", "config/vocabulario.toml", workbook.WORKBOOK_V2.as_posix()])
    assert acuerdo.main(["--root", str(tmp_path)]) == 0
    relativa = acuerdo.SALIDA_REPORTE
    assert (tmp_path / relativa).read_bytes() == (RAIZ_REPO / relativa).read_bytes()


def test_acuerdo_entre_dos_workbooks_anotados(repo_workbook: Path) -> None:
    ruta_a = repo_workbook / workbook.WORKBOOK_V2
    ruta_b = repo_workbook / "data/metadata/anotacion_v2_evaluador2.xlsx"
    ruta_b.write_bytes(ruta_a.read_bytes())
    ids = [f["id_global"] for f in filas_hoja(ruta_a, workbook.HOJA_ANOTACION)][:6]
    comunes = {"color_dominante": "rojo", "contraste": "alto", "colores": "rojo; crema"}
    composiciones = ["bandas", "paneles", "bandas", "reticula", "bandas", "paneles"]
    anotar_v2(ruta_a, {
        id_: {**comunes, "composicion": c, "borde": "si" if i % 2 else "no"}
        for i, (id_, c) in enumerate(zip(ids, composiciones))
    } | {ids[5]: {"composicion": "paneles", "borde": "no_visible", "contraste": "altisimo", "color_dominante": "rojo"}})
    anotar_v2(ruta_b, {
        id_: {**comunes, "composicion": c, "borde": "si" if i % 2 else "no"}
        for i, (id_, c) in enumerate(zip(ids[:5], composiciones))
    } | {ids[5]: {"composicion": "paneles", "borde": "si", "contraste": "alto", "color_dominante": "rojo"}})
    # Una pieza anotada solo por A no se compara.
    extra = [f["id_global"] for f in filas_hoja(ruta_a, workbook.HOJA_FUERA_NUCLEO)][0]
    anotar_v2(ruta_a, {extra: {"composicion": "bandas"}})

    assert acuerdo.main(["--root", str(repo_workbook)]) == 0
    reporte = (repo_workbook / acuerdo.SALIDA_REPORTE).read_text(encoding="utf-8")
    fila = lambda atributo: next(l for l in reporte.splitlines() if l.startswith("| ") and f"| `{atributo}` |" in l)  # noqa: E731
    assert "- Piezas anotadas por los dos (emparejadas por `id_global`): 6." in reporte
    assert fila("composicion").endswith("| 6 | 0 | 1.000 | 0.389 | 1.000 | si |")
    # Borde: la pieza marcada no_visible por A sale del calculo y se cuenta aparte.
    assert "| `borde` | 0 | 0 | 1 |" in reporte
    assert fila("borde").split("|")[4].strip() == "5"
    # Color dominante constante en los dos: no estimable, nunca 1 ni 0.
    assert "no estimable: una sola categoria observada (rojo)" in fila("color_dominante")
    # Valor fuera del vocabulario: se informa y la pieza sale del atributo.
    assert "| `contraste` | 0 | 1 | 0 |" in reporte and "altisimo" in reporte
    # Correccion vacia en los dos: se compara el valor vigente, que es el mismo.
    assert fila("estado_imagen_corregido").split("|")[4].strip() == "6"
