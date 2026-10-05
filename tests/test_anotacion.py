"""Pruebas del vocabulario controlado y del flujo de anotacion manual."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.anotacion import aplicar_anotaciones, workbook
from src.consolidado import construir_consolidado as cc
from src.utils.config import RAIZ_REPO
from src.utils.vocabulario import leer_vocabulario

INSUMOS_CONSOLIDADO = [
    "config/corpus.toml",
    "config/vocabulario.toml",
    "data/processed/corpus_met_textiles_andinos_v1_principal.csv",
    "data/processed/corpus_met_textiles_andinos_v1_complementario.csv",
    "data/processed/corpus_met_textiles_andinos_v1_inventario_base.csv",
    "data/metadata/met_corpus_principal_v2_revisado.csv",
    "data/metadata/met_corpus_secundario_v2_revisado.csv",
    "data/metadata/cma_corpus_principal_revisado.csv",
    "data/metadata/cma_corpus_secundario_revisado.csv",
    "data/metadata/cma_andes_textiles_candidates.csv",
    "data/metadata/imagenes_manifiesto_met_cma_v1.csv",
    cc.SALIDA_CONSOLIDADO,
]


# --------------------------------------------------------------------------- vocabulario


def test_vocabulario_recodifica_etiquetas_y_ordena() -> None:
    material = leer_vocabulario()["material_normalizado"]
    assert material.recodificar("fibra de camélido; algodón") == "algodon; fibra_de_camelido"
    assert material.invalidos("algodon; lino") == ["lino"]


def test_codigos_del_vocabulario_son_ascii_sin_espacios() -> None:
    for campo in leer_vocabulario().values():
        for codigo in campo.codigos:
            assert codigo.isascii() and " " not in codigo and codigo == codigo.lower(), (campo.nombre, codigo)


# --------------------------------------------------------------------------- workbook


@pytest.fixture()
def repo(tmp_path: Path) -> Path:
    for relativa in INSUMOS_CONSOLIDADO:
        destino = tmp_path / relativa
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(RAIZ_REPO / relativa, destino)
    assert workbook.main(["--root", str(tmp_path)]) == 0
    return tmp_path


def anotar(raiz: Path, id_global: str, valores: dict[str, str]) -> None:
    ruta = raiz / workbook.WORKBOOK
    libro = load_workbook(ruta)
    hoja = libro[workbook.HOJA_ANOTACION]
    encabezados = [celda.value for celda in hoja[1]]
    fila = next(f for f in hoja.iter_rows(min_row=2) if f[0].value == id_global)
    for columna, valor in valores.items():
        fila[encabezados.index(columna)].value = valor
    libro.save(ruta)


def leer_anotaciones(raiz: Path) -> dict[str, dict[str, str]]:
    with (raiz / cc.ANOTACIONES).open(newline="", encoding="utf-8") as archivo:
        return {fila["id_global"]: fila for fila in csv.DictReader(archivo)}


def test_workbook_tiene_un_registro_por_fila_y_listas(repo: Path) -> None:
    libro = load_workbook(repo / workbook.WORKBOOK)
    hoja = libro[workbook.HOJA_ANOTACION]
    assert hoja.max_row - 1 == 288
    assert {"guia", "vocabulario"} <= set(libro.sheetnames)
    assert any(dv.type == "list" for dv in hoja.data_validations.dataValidation)


def test_anotacion_valida_se_exporta_y_recodifica(repo: Path) -> None:
    anotar(repo, "MET:307449", {"color_dominante": "marrón", "motivos": "zoomorfo; geométrico", "anotador": "HS"})

    assert aplicar_anotaciones.main(["--root", str(repo)]) == 0
    anotacion = leer_anotaciones(repo)["MET:307449"]
    assert anotacion["color_dominante"] == "marron"
    assert anotacion["motivos"] == "geometrico; zoomorfo"
    assert anotacion["estado_anotacion"] == "anotado_manual"


def test_correccion_marca_corregido_y_llega_al_consolidado(repo: Path) -> None:
    anotar(repo, "CMA:102924", {"tipo_objeto_corregido": "borde_textil", "material_corregido": "algodon; lana"})
    assert aplicar_anotaciones.main(["--root", str(repo)]) == 0
    assert cc.main(["--root", str(repo)]) == 0

    with (repo / cc.SALIDA_CONSOLIDADO).open(newline="", encoding="utf-8") as archivo:
        fila = next(f for f in csv.DictReader(archivo) if f["id_global"] == "CMA:102924")
    assert fila["tipo_objeto"] == "borde_textil" and fila["origen_tipo_objeto"] == "anotacion"
    assert fila["material_normalizado"] == "algodon; lana"
    assert fila["estado_anotacion"] == "corregido"


def test_valor_fuera_del_vocabulario_detiene_y_no_toca_el_csv(repo: Path) -> None:
    assert aplicar_anotaciones.main(["--root", str(repo)]) == 0
    antes = (repo / cc.ANOTACIONES).read_bytes()
    anotar(repo, "MET:307449", {"contraste": "altisimo", "fecha_anotacion": "05/10/2026"})

    assert aplicar_anotaciones.main(["--root", str(repo)]) == 1
    assert (repo / cc.ANOTACIONES).read_bytes() == antes
    reporte = (repo / aplicar_anotaciones.REPORTE).read_text(encoding="utf-8")
    assert "altisimo" in reporte and "AAAA-MM-DD" in reporte


def test_regenerar_workbook_conserva_lo_anotado(repo: Path) -> None:
    anotar(repo, "MET:307449", {"familia_iconografica": "geometrica", "observaciones_anotacion": "bandas escalonadas"})
    assert workbook.main(["--root", str(repo)]) == 0

    hoja = load_workbook(repo / workbook.WORKBOOK)[workbook.HOJA_ANOTACION]
    encabezados = [celda.value for celda in hoja[1]]
    fila = next(f for f in hoja.iter_rows(min_row=2, values_only=True) if f[0] == "MET:307449")
    assert fila[encabezados.index("familia_iconografica")] == "geometrica"
    assert fila[encabezados.index("observaciones_anotacion")] == "bandas escalonadas"
