"""Pruebas de la ficha de aptitud: medidas, motivos, puntaje, calibracion y workbooks."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from src.imagenes import aptitud, calibrar_aptitud, medidas
from src.imagenes.aptitud import motivo_pertinencia, normalizar, predecir, puntuar_registro
from src.imagenes.aptitud_workbooks import (
    PREFIJO_OBSERVACION,
    discrepancias,
    editada,
    es_decision_automatica,
    intacta,
    motivo_del_revisor,
    motivo_enriquecido,
)
from src.imagenes.calibrar_aptitud import logistica
from src.imagenes.descargar_evaluacion import ENTRADA_AIC, ENTRADA_CMA, ENTRADAS_MET
from src.imagenes.medidas import COLUMNAS_MEDIDAS, etiquetar, medir_imagen
from src.imagenes.motivos import analizar_motivo, describir_motivo
from src.utils.config import RAIZ_REPO, leer_config_aptitud
from src.utils.fechas import anio_inicio
from src.utils.metricas import auc
from tests.insumos import copiar_insumos

CONFIG = leer_config_aptitud(RAIZ_REPO)
FONDO = (128, 128, 128)


# ------------------------------------------------------------------ medidas con imagenes sinteticas

def guardar(tmp_path: Path, matriz: np.ndarray, nombre: str = "img.png") -> Path:
    ruta = tmp_path / nombre
    Image.fromarray(matriz.astype(np.uint8)).save(ruta)
    return ruta


def tablero(alto: int, ancho: int, celda: int) -> np.ndarray:
    """Reticula de modulos de cuatro colores, como una banda de tocapus."""
    colores = np.array([[200, 30, 30], [240, 200, 40], [20, 60, 160], [245, 240, 225]])
    y, x = np.mgrid[0:alto, 0:ancho]
    return colores[((y // celda) * 3 + (x // celda)) % 4]


def lienzo(alto: int = 1200, ancho: int = 1600) -> np.ndarray:
    return np.full((alto, ancho, 3), FONDO, dtype=np.int64)


def medidas_imagen(tmp_path: Path, matriz: np.ndarray, lado_original: int = 4000) -> dict:
    escala = lado_original / max(matriz.shape[:2])
    ruta = guardar(tmp_path, matriz)
    return medir_imagen(ruta, round(matriz.shape[1] * escala), round(matriz.shape[0] * escala), CONFIG["segmentacion"])


def test_normalizar_admite_anclas_invertidas_y_logaritmicas() -> None:
    assert normalizar(0.5, {"bajo": 0.0, "alto": 1.0}) == 0.5
    assert normalizar(0.02, {"bajo": 0.3, "alto": 0.02}) == 1.0
    assert normalizar(5.0, {"bajo": 0.0, "alto": 1.0}) == 1.0
    assert normalizar(1000, {"bajo": 100, "alto": 10000, "escala": "log"}) == pytest.approx(0.5)


def test_etiquetar_cuenta_componentes_en_vecindad_4() -> None:
    mascara = np.zeros((10, 10), dtype=bool)
    mascara[1:4, 1:4] = True
    mascara[6:9, 6:9] = True
    mascara[4, 4] = True  # toca la primera solo en diagonal: es otra componente
    etiquetas, n = etiquetar(mascara)
    assert n == 3
    assert len(np.unique(etiquetas[1:4, 1:4])) == 1


def test_textil_con_reticula_de_colores(tmp_path: Path) -> None:
    imagen = lienzo()
    imagen[200:1000, 300:1300] = tablero(800, 1000, 50)
    medidas = medidas_imagen(tmp_path, imagen)
    assert medidas["n_fragmentos"] == 1
    assert medidas["fraccion_lagunas"] < 0.02
    assert medidas["regularidad"] > CONFIG["estructura"]["regularidad_minima"]
    assert medidas["colores_distintos"] >= 3  # rojo, amarillo, azul y crudo: claramente distintos
    assert medidas["bordes_cromaticos"] > 0.1
    assert medidas["toca_encuadre"] == 0.0


def test_fragmentos_sueltos_y_lagunas_se_detectan(tmp_path: Path) -> None:
    imagen = lienzo()
    imagen[150:550, 150:650] = tablero(400, 500, 50)
    imagen[650:1050, 900:1400] = tablero(400, 500, 50)
    medidas = medidas_imagen(tmp_path, imagen)
    assert medidas["n_fragmentos"] == 2
    assert medidas["fraccion_mayor"] < 0.6

    con_laguna = lienzo()
    con_laguna[200:1000, 300:1300] = tablero(800, 1000, 50)
    con_laguna[450:750, 600:1000] = FONDO
    assert medidas_imagen(tmp_path, con_laguna)["fraccion_lagunas"] > 0.1


def test_tela_lisa_no_tiene_colores_distintos_ni_bordes(tmp_path: Path) -> None:
    imagen = lienzo()
    imagen[200:1000, 300:1300] = (180, 150, 110)
    medidas = medidas_imagen(tmp_path, imagen)
    assert medidas["colores_distintos"] <= 1
    assert medidas["n_tonos"] <= 1
    assert medidas["bordes_cromaticos"] < 0.01


def test_banda_alargada(tmp_path: Path) -> None:
    imagen = lienzo()
    imagen[550:650, 100:1500] = tablero(100, 1400, 50)
    assert medidas_imagen(tmp_path, imagen)["elongacion"] > 8


# ------------------------------------------------------------------ motivos del revisor

@pytest.mark.parametrize(
    ("motivo", "esperado"),
    [
        ("Textil compelto con colores y formas visibles. La mejor imagen para el proyecto.",
         {"completitud": "completo", "colores": "si", "formas": "si", "referencia": True}),
        ("Textil incompelto con colores y formas visibles", {"completitud": "incompleto", "colores": "si", "formas": "si"}),
        ("Textil incompleto sin formas ni colores visibles", {"completitud": "incompleto", "colores": "no", "formas": "no"}),
        ("Textil con colores visibles y no formas", {"completitud": "", "colores": "si", "formas": "no"}),
        ("Textil incompleto con formas y colores poco visibles", {"colores": "parcial", "formas": "parcial"}),
        ("textil incompleto con figuras visibles", {"formas": "si", "colores": ""}),
        ("Iconografía completa y colores visibles", {"formas": "si", "colores": "si"}),
        ("Textil con Plumas no es valido", {"pertinencia": "plumaria"}),
        ("Son muñecos ", {"pertinencia": "muneco"}),
        ("No es textil andino Inca o Anterior, es de 1900 y es de la Selva", {"pertinencia": "posterior_inca"}),
        ("herramienta_textil, no es un textil andino", {"pertinencia": "herramienta_textil"}),
        ("fotografia_monocroma Textil completo con forma visible", {"monocroma": True, "completitud": "completo", "formas": "si"}),
    ],
)
def test_analizar_motivo(motivo: str, esperado: dict) -> None:
    resultado = analizar_motivo(motivo)
    for clave, valor in esperado.items():
        assert resultado[clave] == valor, clave


def test_describir_motivo_normaliza() -> None:
    texto = describir_motivo(analizar_motivo("Textil compelto con colores y formas visibles. La mejor imagen para el proyecto."))
    assert texto == "Textil completo; colores visibles; formas visibles; imagen de referencia para el estudio"


# ------------------------------------------------------------------ puntaje con un modelo sintetico

def fila_medidas(**valores: object) -> dict[str, str]:
    base = {campo: "0" for campo in COLUMNAS_MEDIDAS}
    base.update({"id_global": "AIC:1", "fuente": "AIC", "id_fuente": "1", "estado_imagen": "ok", "sha256_copia": "",
                 "encuadre_lleno": "no", "terminos_fragmento": "", "lado_textil_px": "2000", "nitidez": "1.5", "nitidez_comun": "1.5",
                 "croma_p90": "40", "colorido": "40", "n_tonos": "3", "colores_distintos": "4", "compacidad": "0.9"})
    base.update({clave: str(valor) for clave, valor in valores.items()})
    return base


def modelo_sintetico(sesgo: float) -> dict:
    """Cada atributo vale sigma(sesgo): sesgo alto da puntaje alto."""
    atributos = {nombre: {"medidas": ["nitidez"], "pesos": [sesgo, 0.0], "n": 1} for nombre in CONFIG["atributos"]}
    return {
        "estandarizacion": {"nitidez": [0.0, 1.0]},
        "atributos": atributos,
        "puntaje": {"componentes": CONFIG["puntaje"]["componentes"], "pesos": [0.0, 1.0, 1.0, 1.0, 1.0]},
        "umbrales": {"referencia": 90.0, "principal": 60.0, "descartar": 40.0},
    }


@pytest.mark.parametrize(
    ("sesgo", "banda", "referencia"),
    [(5.0, "principal", "si"), (1.0, "principal", "no"), (0.0, "revisar", "no"), (-3.0, "descartar", "no")],
    ids=["referencia", "principal", "revisar", "descartar"],
)
def test_bandas_por_puntaje(sesgo: float, banda: str, referencia: str) -> None:
    resultado = puntuar_registro(fila_medidas(), CONFIG, modelo_sintetico(sesgo))
    assert (resultado["banda_sugerida"], resultado["referencia_sugerida"]) == (banda, referencia)
    assert resultado["justificacion"].startswith("Fotografía")


@pytest.mark.parametrize(
    ("valores", "pertinencia", "banda", "motivo"),
    [
        ({"estado_imagen": "sin_imagen"}, "", "descartar", "sin_imagen"),
        ({}, "posterior_inca", "descartar", "posterior_inca"),
        ({"lado_textil_px": "200"}, "", "descartar", "resolucion_insuficiente"),
        ({"croma_p90": "1"}, "", "secundario", "fotografia_monocroma"),
    ],
    ids=["sin_imagen", "pertinencia", "resolucion", "monocroma"],
)
def test_puertas_antes_del_puntaje(valores: dict, pertinencia: str, banda: str, motivo: str) -> None:
    resultado = puntuar_registro(fila_medidas(**valores), CONFIG, modelo_sintetico(5.0), pertinencia)
    assert (resultado["banda_sugerida"], resultado["motivo_banda"]) == (banda, motivo)
    assert resultado["referencia_sugerida"] == "no"


def test_monocroma_nunca_tiene_colores_identificables() -> None:
    resultado = puntuar_registro(fila_medidas(croma_p90="1"), CONFIG, modelo_sintetico(5.0))
    assert resultado["colores"] == "no_identificables"
    assert resultado["justificacion"].startswith("Fotografía en blanco y negro")


@pytest.mark.parametrize(
    ("titulo", "material", "fecha", "motivo"),
    [
        ("Feathered Tabard", "cotton (textile); feather", 1470, "plumaria"),
        ("Half of a Tabard (Open-sided Tunic)", "feather; cotton (fiber)", 1250, "plumaria"),
        ("Dolls", "cotton (fiber)", 1950, "muneco"),
        ("Chief's Poncho", "wool (textile)", 1850, "posterior_inca"),
        ("Tunic with Tocapu", "cotton; camelid fiber", 1450, ""),
        ("Fragment (Band)", "cotton", None, ""),
        ("Loincloth with Feather(?) Motifs", "cotton, camelid fiber", 1000, ""),  # iconografia, no plumas
        ("Featherstitched Border", "", 1450, ""),  # palabra completa: featherstitched no es plumaria
    ],
    ids=["plumaria_titulo", "plumaria_material", "muneco", "posterior", "inca", "sin_fecha", "motivo_pluma", "palabra"],
)
def test_reglas_de_pertinencia(titulo: str, material: str, fecha: int | None, motivo: str) -> None:
    metadatos = {"titulo": titulo, "material": material, "fecha_inicio": fecha}
    assert motivo_pertinencia(metadatos, CONFIG["pertinencia"]) == motivo


@pytest.mark.parametrize(
    ("texto", "anio"),
    [
        ("7th–9th century", 600), ("5th–2nd century BCE", -500), ("2nd century BCE", -200),
        ("late 16th–17th century", 1566), ("c. 1400–1532", 1400), ("400–200 BCE", -400),
        ("100 BCE–700 CE", -100), ("1600s–1700s", 1600), ("early 1900s", 1900), ("mid-1800s", 1833),
        ("1460–1632 (radiocarbon date, 95% probability)", 1460), ("1000–1460s", 1000), ("", None),
    ],
)
def test_anio_inicio(texto: str, anio: int | None) -> None:
    assert anio_inicio(texto) == anio


def test_nitidez_comun_no_depende_de_la_resolucion(tmp_path: Path) -> None:
    imagen = lienzo()
    imagen[200:1000, 300:1300] = tablero(800, 1000, 50)
    grande = medidas_imagen(tmp_path, imagen)
    reducida = np.asarray(Image.fromarray(imagen.astype(np.uint8)).resize((800, 600), Image.Resampling.BOX))
    pequena = medidas_imagen(tmp_path, reducida)
    assert pequena["nitidez"] > 1.3 * grande["nitidez"]  # la nitidez de la copia sube al reducir
    assert pequena["nitidez_comun"] == pytest.approx(grande["nitidez_comun"], rel=0.15)


# ------------------------------------------------------------------ calibracion

def test_logistica_separa_datos_separables() -> None:
    rng = np.random.default_rng(0)
    x = rng.normal(size=(200, 2))
    y = (x[:, 0] > 0).astype(float)
    w = logistica(x, y, l2=1.0)
    p = predecir(x, w)
    assert auc(p[y == 1].tolist(), p[y == 0].tolist()) > 0.99
    assert w[1] > abs(w[2])


def test_auc() -> None:
    assert auc([3, 4, 5], [1, 2]) == 1.0
    assert auc([1, 2], [1, 2]) == 0.5


# ------------------------------------------------------------------ workbooks

def lector(fila: dict) -> callable:
    return lambda campo: fila.get(campo, "")


AUTOMATICA = {
    "origen_curacion": "aptitud_automatica", "revisor": "auto", "motivo_revision": "fotografia_monocroma",
    "corpus_final": "secundario", "corpus_actual": "principal", "decision_revision": "mover",
    "observaciones_revision": f"{PREFIJO_OBSERVACION} Fotografia en blanco y negro: el color no es identificable. Puntaje 50.",
}


def test_fila_automatica_intacta_y_editada() -> None:
    assert intacta(lector(AUTOMATICA)) and es_decision_automatica(lector(AUTOMATICA))
    # El motivo enriquecido no rompe la deteccion: se basa en la observacion.
    enriquecida = {**AUTOMATICA, "motivo_revision": "Decisión automática: Fotografía monocroma. Análisis de imagen: ..."}
    assert intacta(lector(enriquecida))
    corregida = {**AUTOMATICA, "corpus_final": "principal"}  # el revisor la paso a principal sin cambiar revisor
    assert editada(lector(corregida)) and not es_decision_automatica(lector(corregida))
    herramienta = {"origen_curacion": "descarte_automatico_aic", "revisor": "auto", "motivo_revision": "herramienta_textil"}
    assert es_decision_automatica(lector(herramienta))


def test_motivo_enriquecido_conserva_el_original_y_no_se_duplica() -> None:
    original = "Textil compelto con colores y formas visibles"
    enriquecido = motivo_enriquecido(original, "Fotografía buena (nitidez 2,00).", automatica=False)
    assert enriquecido == "Revisor: Textil completo; colores visibles; formas visibles. Análisis de imagen: Fotografía buena (nitidez 2,00)."
    fila = {"motivo_revision": enriquecido, "motivo_generado": enriquecido, "motivo_revision_original": original}
    assert motivo_del_revisor(lector(fila)) == original  # segunda ejecucion: parte del original
    editado = {**fila, "motivo_revision": "Textil incompleto"}
    assert motivo_del_revisor(lector(editado)) == "Textil incompleto"  # el revisor lo edito: es su nuevo original


def test_discrepancias() -> None:
    registro = {"prob_colores": "0.95", "prob_formas": "0.05", "prob_completitud": "0.92", "prob_referencia": "0.9",
                "referencia_sugerida": "si", "banda_sugerida": "principal", "puntaje_aptitud": "88"}
    notas = discrepancias("Textil con colores y formas visibles", "principal", registro, 0.85)
    assert "posible textil completo no anotado" in notas
    assert "anotado con formas visibles, pero la imagen sugiere lo contrario" in notas
    assert "posible imagen de referencia" in notas
    assert "colores" not in notas.split("formas")[0]  # colores coincide: no se marca


# ------------------------------------------------------------------ regeneracion

def test_modelo_y_fichas_se_regeneran_igual(tmp_path: Path) -> None:
    copiar_insumos(tmp_path, ("config/aptitud.toml", medidas.MEDIDAS, "data/metadata/aic_revision_manual_v2.xlsx",
                              "data/metadata/met_revision_manual_v2.xlsx", "data/metadata/cma_revision_manual_v2.xlsx",
                              aptitud.INVENTARIO_MET, *ENTRADAS_MET, ENTRADA_CMA, ENTRADA_AIC))
    assert calibrar_aptitud.main(["--root", str(tmp_path)]) == 0
    assert aptitud.main(["--root", str(tmp_path)]) == 0
    for relativa in (calibrar_aptitud.MODELO, aptitud.SALIDA_APTITUD):
        generado = (tmp_path / relativa).read_bytes()
        assert b"\r\n" not in generado
        assert generado == (RAIZ_REPO / relativa).read_bytes(), relativa


# ------------------------------------------------------------------ validacion

def test_muestra_de_validacion_estratificada_y_reproducible() -> None:
    from src.imagenes.validacion import MINIMO_POR_BANDA, TAMANO_POR_FUENTE, muestrear

    candidatos = [{"id_global": f"{f}:{b}-{n}", "fuente": f, "banda": b}
                  for f in ("MET", "CMA") for b, total in (("principal", 100), ("revisar", 25), ("descartar", 6))
                  for n in range(total)]
    muestra = muestrear(candidatos)
    assert muestra == muestrear(candidatos)  # semilla fija
    for fuente, tamano in TAMANO_POR_FUENTE.items():
        de_fuente = [m for m in muestra if m["fuente"] == fuente]
        assert len(de_fuente) == tamano
        assert all(sum(m["banda"] == b for m in de_fuente) >= MINIMO_POR_BANDA for b in ("principal", "revisar", "descartar"))
    assert sorted(int(m["orden"]) for m in muestra) == list(range(1, len(muestra) + 1))


def test_referencia_por_completitud() -> None:
    completo = analizar_motivo("Textil compelto con colores y formas visibles")
    assert completo["referencia"] and not completo["referencia_marcada"]
    assert not analizar_motivo("Textil incompleto con colores y formas visibles")["referencia"]
    assert not analizar_motivo("Textil completo sin colores ni formas visibles")["referencia"]
    marcada = analizar_motivo("Textil incompleto con colores y formas visibles. La mejor imagen para el proyecto.")
    assert marcada["referencia"] and marcada["referencia_marcada"]


@pytest.mark.parametrize(
    ("decision", "motivo", "contradice"),
    [
        ("descartados", "Textil incompleto con colores visibles y figuras visibles", True),
        ("principal", "textil sin formas ni colores visibles", True),
        ("principal", "textil incompleto con figuras visibles", False),
        ("secundario", "Textil con colores visibles", False),
        ("descartados", "Textil con Plumas no es valido, con colores y formas visibles", False),
    ],
)
def test_contradice_motivo(decision: str, motivo: str, contradice: bool) -> None:
    from src.imagenes.traslado import contradice_motivo

    assert contradice_motivo(decision, motivo) == contradice


def test_traslado_de_la_revision_v2(tmp_path: Path) -> None:
    """Cambio, confirmacion con motivo nuevo y contradiccion; volver a trasladar no cambia nada."""
    from openpyxl import Workbook, load_workbook

    from src.imagenes.traslado import trasladar_anotaciones
    from src.utils.fuentes import FUENTES

    copiar_insumos(tmp_path, (FUENTES["MET"].workbook, FUENTES["CMA"].workbook))
    ruta, hoja = FUENTES["MET"].workbook, FUENTES["MET"].hoja
    filas = list(load_workbook(tmp_path / ruta, read_only=True)[hoja].iter_rows(values_only=True))
    indice = {n: i for i, n in enumerate(filas[0])}
    vigentes = {f"MET:{f[indice['id_fuente']]}": f[indice["corpus_final"]] for f in filas[1:]}
    secundario = next(i for i, d in vigentes.items() if d == "secundario")
    principal = next(i for i, d in vigentes.items() if d == "principal")
    otro = next(i for i, d in vigentes.items() if d == "secundario" and i != secundario)

    libro = Workbook()
    libro.active.title = "revision"
    libro.active.append(["orden", "id_global", "corpus_final", "motivo_revision"])
    libro.active.append(["1", secundario, "principal", "Textil incompleto con colores y formas visibles"])
    libro.active.append(["2", principal, "principal", "Textil completo con colores y formas visibles"])
    libro.active.append(["3", otro, "descartados", "Textil con colores y formas visibles"])  # contradice su motivo
    libro.save(tmp_path / "revision.xlsx")

    conteo = trasladar_anotaciones(tmp_path, "revision.xlsx", "revision", "2026-10-20", "revision_criterio_v2",
                                   "Revisión con el criterio v2", registrar_confirmaciones=True)
    assert (conteo["cambiadas"], conteo["confirmadas"], conteo["revisar"]) == (1, 1, 1)
    filas = list(load_workbook(tmp_path / ruta, read_only=True)[hoja].iter_rows(values_only=True))
    nuevas = {f"MET:{f[indice['id_fuente']]}": f for f in filas[1:]}
    assert nuevas[secundario][indice["corpus_final"]] == "principal"
    assert "Decisión anterior: secundario" in nuevas[secundario][indice["observaciones_revision"]]
    assert nuevas[principal][indice["motivo_revision"]] == "Textil completo con colores y formas visibles"
    assert nuevas[otro][indice["corpus_final"]] == "revisar"

    repetido = trasladar_anotaciones(tmp_path, "revision.xlsx", "revision", "2026-10-20", "revision_criterio_v2",
                                     "Revisión con el criterio v2", registrar_confirmaciones=True)
    assert sum(repetido.values()) == 0


@pytest.mark.parametrize(
    ("titulo", "marcada"),
    [
        ("Border Fragment, Figure", True),
        ("Male Figure with Modern Additions", True),
        ("Embroidered Fragment with Figure", False),   # figura representada en el textil
        ("Panel with a Seated Figure", False),
        ("Tunic with Frontal Figures", False),         # plural: siempre representadas
        ("Figured Cloth", False),
    ],
)
def test_posible_figurilla_es_prudente(titulo: str, marcada: bool) -> None:
    from src.imagenes.aptitud import posible_figurilla

    assert posible_figurilla(titulo, CONFIG["pertinencia"]) == marcada


def test_figurilla_va_a_revisar_y_no_a_descartar() -> None:
    resultado = puntuar_registro(fila_medidas(), CONFIG, modelo_sintetico(5.0), "posible_figurilla")
    assert (resultado["banda_sugerida"], resultado["motivo_banda"]) == ("revisar", "posible_figurilla")
