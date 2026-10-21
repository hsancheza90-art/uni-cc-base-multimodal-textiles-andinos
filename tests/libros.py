"""Contenido comparable de un workbook, para pruebas de no regresion.

openpyxl guarda la fecha de creacion dentro del xlsx, asi que dos workbooks iguales no
coinciden byte a byte; se compara lo que ve quien lo abre: hojas, celdas, enlaces, listas
desplegables, anchos, paneles fijos y filtros.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def contenido_libro(ruta: Path) -> dict[str, Any]:
    libro = load_workbook(ruta)
    contenido: dict[str, Any] = {"hojas": libro.sheetnames}
    for hoja in libro.worksheets:
        contenido[hoja.title] = {
            "celdas": [[celda.value for celda in fila] for fila in hoja.iter_rows()],
            "enlaces": [(c.coordinate, c.hyperlink.target) for fila in hoja.iter_rows() for c in fila if c.hyperlink],
            "rellenos": [c.fill.fgColor.rgb for c in hoja[1]] if hoja.max_row else [],
            "validaciones": sorted(
                (str(v.sqref), v.type, v.formula1, v.prompt, v.showErrorMessage, v.showInputMessage)
                for v in hoja.data_validations.dataValidation
            ),
            "anchos": {letra: dim.width for letra, dim in sorted(hoja.column_dimensions.items())},
            "paneles": hoja.freeze_panes,
            "filtro": hoja.auto_filter.ref,
        }
    return contenido
